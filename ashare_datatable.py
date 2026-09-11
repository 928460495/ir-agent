# -*- coding: utf-8 -*-
"""可选：在「1_假设」页追加 Excel 原生二维模拟运算表（Data Table）。

必须在验算通过之后调用——formulas 引擎无法解析 DataTableFormula，
一旦写入，整个工作簿都不能再被引擎验算。

Excel 约束：模拟运算表的行/列输入单元格必须与表位于同一工作表，
因此三张表建在「1_假设」页（Δ 输入格与永续 g 都在该页）。
"""
from __future__ import annotations
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.utils import get_column_letter as CL

NAVY="1F4E79"; LBLUE="D9E1F2"; GREY="F2F2F2"
F_HDR=Font(color="FFFFFF",bold=True,size=11); F_COL=Font(bold=True,size=10)
P_N=PatternFill("solid",fgColor=NAVY); P_B=PatternFill("solid",fgColor=LBLUE); P_G=PatternFill("solid",fgColor=GREY)
CNY="¥#,##0.00"; PCT2="0.00%"

def _rowmap(ws):
    m={}
    for r in range(1,ws.max_row+1):
        v=ws.cell(r,1).value
        if isinstance(v,str) and v.strip(): m.setdefault(v.strip(),r)
    return m
def _find(m,p):
    if p in m: return m[p]
    hits=[r for k,r in m.items() if k.startswith(p)]
    if len(hits)==1: return hits[0]
    raise KeyError(p)

def add(path: str) -> list[str]:
    wb=load_workbook(path)
    a=wb["1_假设"]; am=_rowmap(a)
    d=wb["5_DCF"]; dm=_rowmap(d)
    ps_row=_find(dm,"每股内在价值")
    R_DG=_find(am,"营收增速 Δ"); R_DM=_find(am,"毛利率 Δ"); R_G=_find(am,"永续增长率 g")
    r=a.max_row+3
    a.cell(r,1,"六、Excel 原生模拟运算表（What-If / Data Table）").font=F_HDR
    for c in range(1,8): a.cell(r,c).fill=P_N
    r+=1
    a.cell(r,1,"由 Excel 自身重算，非本模板公式。若显示为 0，请按 F9 或将「公式→计算选项」设为「自动」。").font=Font(size=9,color="C00000"); r+=1
    a.cell(r,1,"⚠ 本页含 DataTable，工作簿将无法再被 formulas 引擎验算——「5_DCF」的引擎版敏感性表仍可作交叉核对。").font=Font(size=9,color="C00000"); r+=2
    made=[]
    specs=[("模拟表 1：营收增速 Δ（列）× 永续 g（行）",[-.10,-.05,0.0,.05,.10],
            [.020,.025,.030,.035,.040], f"$B${R_G}", f"$B${R_DG}", "营收增速 Δ \\ 永续 g"),
           ("模拟表 2：毛利率 Δ（列）× 永续 g（行）",[-.04,-.02,0.0,.02,.04],
            [.020,.025,.030,.035,.040], f"$B${R_G}", f"$B${R_DM}", "毛利率 Δ \\ 永续 g"),
           ("模拟表 3：营收增速 Δ（列）× 毛利率 Δ（行）",[-.10,-.05,0.0,.05,.10],
            [-.04,-.02,0.0,.02,.04], f"$B${R_DM}", f"$B${R_DG}", "营收增速 Δ \\ 毛利率 Δ")]
    for title,lefts,tops,r1,r2,corner_lbl in specs:
        a.cell(r,1,title).font=Font(bold=True,size=11,color=NAVY); r+=1
        a.cell(r,1,f"行输入单元格 {r1}　｜　列输入单元格 {r2}").font=Font(size=9,color="808080"); r+=1
        c0=r
        cc=a.cell(r,2,f"='5_DCF'!B{ps_row}")     # 左上角 = 输出公式
        cc.font=Font(bold=True,size=10); cc.number_format=CNY
        a.cell(r,1,corner_lbl).font=F_COL; a.cell(r,1).fill=P_B
        for j,t in enumerate(tops):
            x=a.cell(r,3+j,t); x.font=F_COL; x.fill=P_B; x.number_format=PCT2; x.alignment=Alignment(horizontal="center")
        r+=1
        first=r
        for lv in lefts:
            x=a.cell(r,2,lv); x.font=F_COL; x.fill=P_G; x.number_format=PCT2
            r+=1
        last=r-1
        ref=f"C{first}:{CL(2+len(tops))}{last}"
        for rr in range(first,last+1):
            for cc_ in range(3,3+len(tops)):
                cell=a.cell(rr,cc_)
                cell.value=DataTableFormula(ref=ref, dt2D=True, dtr=False, r1=r1.replace("$",""), r2=r2.replace("$",""))
                cell.number_format=CNY
        made.append(f"{title}  →  {ref}")
        r+=2
    wb.save(path)
    return made
