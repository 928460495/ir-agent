"""DCF 工作簿的「9_相对估值」页。用 ashare-dcf 的生益科技固件建真工作簿。"""

from __future__ import annotations

import copy
import pickle
import shutil
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from ir_agent import dcf

D = Decimal
FIX = dcf.DEFAULT_DIR / "tests" / "fixtures" / "cal_600183.pkl"

try:
    dcf._load(dcf.DEFAULT_DIR)
    import ashare_build as B
    import ashare_mc as MC
except Exception:                                           # noqa: BLE001
    B = MC = None

pytestmark = pytest.mark.skipif(B is None or not FIX.exists(),
                                reason="缺 ashare-dcf 或固件")


@pytest.fixture(scope="module")
def base_wb(tmp_path_factory):
    with open(FIX, "rb") as f:
        cal = pickle.load(f)
    p = tmp_path_factory.mktemp("rv") / "base.xlsx"
    B.build(copy.deepcopy(cal), str(p))
    MC.run(str(p), hist=dcf._mc_hist(cal), quiet=True)
    return p


@pytest.fixture
def wb(base_wb, tmp_path):
    p = tmp_path / "w.xlsx"
    shutil.copy(base_wb, p)
    return str(p)


def _res():
    from tests.test_relval import _peers_pe, make_result
    r = make_result()
    r.comps_peers = _peers_pe(10, 20, 30, 40, 50, 60)
    return r


DCF_OK = SimpleNamespace(scenario_ps={"bear": D("40"), "base": D("56.55"), "bull": D("80")},
                         mc={"p5": D("21.92"), "p95": D("137.37")})
DCF_BARE = SimpleNamespace(scenario_ps={}, mc={})


def _text(path):
    ws = load_workbook(path)["9_相对估值"]
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_新页追加在蒙特卡洛之后(wb):
    from ir_agent.relval_sheet import SHEET, write_sheet
    write_sheet(wb, _res(), DCF_OK)
    assert load_workbook(wb).sheetnames[-2:] == ["8_蒙特卡洛", SHEET]


def test_蒙特卡洛图表不丢(wb):
    from ir_agent.relval_sheet import write_sheet
    before = len(load_workbook(wb)["8_蒙特卡洛"]._charts)
    assert before > 0
    write_sheet(wb, _res(), DCF_OK)
    assert len(load_workbook(wb)["8_蒙特卡洛"]._charts) == before


def test_隐含价值是公式(wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(wb, _res(), DCF_OK)
    ws = load_workbook(wb)["9_相对估值"]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    # 三个倍数 × P25 / 中位数 / P75
    assert len(formulas) >= 9


def test_足球场与两张散点图都在(wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(wb, _res(), DCF_OK)
    assert len(load_workbook(wb)["9_相对估值"]._charts) == 3


def test_足球场含现价与各方法(wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(wb, _res(), DCF_OK)
    t = _text(wb)
    for label in ("DCF 三档", "蒙特卡洛 P5–P95", "可比 PE(TTM)", "回归 PB 对 ROE", "现价"):
        assert label in t


def test_缺DCF三档与蒙特卡洛时不画且说明(wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(wb, _res(), DCF_BARE)
    t = _text(wb)
    assert "DCF 三档未生成" in t and "蒙特卡洛未生成" in t


def test_不可用回归在页面写明原因(wb):
    from ir_agent.relval_sheet import write_sheet
    from tests.test_relval import make_result
    write_sheet(wb, make_result(pe_unusable=True), DCF_OK)
    assert "斜率不显著" in _text(wb)


def test_重写不重复建页(wb):
    from ir_agent.relval_sheet import SHEET, write_sheet
    write_sheet(wb, _res(), DCF_OK)
    write_sheet(wb, _res(), DCF_OK)
    assert load_workbook(wb).sheetnames.count(SHEET) == 1


def test_口径说明写明EV含少数股东权益(wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(wb, _res(), DCF_OK)
    assert "少数股东权益" in _text(wb)


# ── 审查意见 1、4 与测试缺口：公式算得对、图引用对、负下端不画、退路口径可见 ──

import re  # noqa: E402


def test_可比法公式算出的就是区间值(wb):
    """只数「=」开头的格子不够：要确认公式引用的是本司指标格，算出来等于区间值。"""
    from ir_agent.relval_sheet import write_sheet
    res = _res()
    write_sheet(wb, res, DCF_OK)
    ws = load_workbook(wb)["9_相对估值"]
    row = next(r for r in ws.iter_rows()
               if r[0].value == "PE(TTM)" and isinstance(r[4].value, str))
    refs = re.findall(r"([A-Z])(\d+)", row[4].value[1:])
    m, np_, sh = (ws[f"{c}{n}"].value for c, n in refs)
    assert m * np_ / sh == pytest.approx(float(res.comps["pe"].lo))


def test_散点图引用的正是样本行(wb):
    from ir_agent.relval_sheet import write_sheet
    res = _res()
    write_sheet(wb, res, DCF_OK)
    ws = load_workbook(wb)["9_相对估值"]
    sc = next(c for c in ws._charts if c.tagname == "scatterChart")
    a, b = map(int, re.findall(r"\$(\d+)", sc.series[0].yVal.numRef.f))
    sample = res.regs["pb_roe"].sample
    assert b - a + 1 == len(sample)
    assert ws[f"A{a}"].value == sample[0][0] and ws[f"A{b}"].value == sample[-1][0]


def test_下端为负的区间不画并说明(wb):
    from ir_agent.relval_sheet import write_sheet
    neg = SimpleNamespace(scenario_ps={"bear": D("-5"), "base": D("10"), "bull": D("30")},
                          mc={"p5": D("21.92"), "p95": D("137.37")})
    write_sheet(wb, _res(), neg)
    ws = load_workbook(wb)["9_相对估值"]
    drawn = [r[0].value for r in ws.iter_rows()
             if isinstance(r[0].value, str) and r[0].value.startswith("DCF 三档")
             and isinstance(r[1].value, (int, float))]
    assert not drawn and "下端为负" in _text(wb)


def test_一年增速的样本在页面标出(wb):
    from ir_agent.relval_sheet import write_sheet
    from tests.test_relval import make_result
    res = make_result()
    res.regs["pe_growth"].flags = {res.regs["pe_growth"].sample[0][0]: "一年增速"}
    write_sheet(wb, res, DCF_OK)
    assert "一年增速" in _text(wb)


def test_写明基本面取最近完整财年(wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(wb, _res(), DCF_OK)
    assert "最近完整财年" in _text(wb)


def test_可比表列出每家的利润口径(wb):
    from ir_agent.relval_sheet import write_sheet
    res = _res()
    res.comps_peers[0].np_basis = "TTM 截至 2026Q1"
    write_sheet(wb, res, DCF_OK)
    assert "利润口径" in _text(wb) and "TTM 截至 2026Q1" in _text(wb)
