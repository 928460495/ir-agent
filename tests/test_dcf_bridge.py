"""接 ashare-dcf 的桥。两个静默出错入口在这里锁住。"""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import pytest

from ir_agent import dcf
from ir_agent.ledger import Fact, FactLedger, Method

D = Decimal
AS_OF = date(2026, 9, 7)
YI = D("100000000")


def _led(**kv) -> FactLedger:
    led = FactLedger()
    for k, v in kv.items():
        led.put(Fact(key=k, value=D(str(v)), unit="元", currency="CNY",
                     period="2025FY", as_of=date(2026, 4, 10),
                     source_id="em_x", method=Method.REPORTED))
    return led


# **用真实结构，不用我想象的结构。** 第一版固件把 stmts 写成字典嵌套，
# 于是 cross_check 也照着字典写 —— 实跑时才发现它是 pandas DataFrame
# （按报告日索引、列为科目名、值为**元**），当场抛
# 「The truth value of a DataFrame is ambiguous」。
#
# 取值一律走 ashare_data.val()（它负责 /1e8 换算），不在这边另写一份。
_D = None
try:
    _D = dcf._load(dcf.DEFAULT_DIR)[0]
except dcf.DcfUnavailable:
    pass

needs_ashare = pytest.mark.skipif(
    _D is None, reason="未找到 ashare-dcf，跳过依赖其真实结构的用例")


def _cal(**items):
    """calibrate 输出片段。**stmts 是 DataFrame，值以元计。**"""
    import pandas as pd
    def frame(cols):
        return pd.DataFrame({k: [v * 1e8] for k, v in cols.items()
                             if v is not None},
                            index=pd.Index(["20251231"], name="报告日"))
    return {"meta": {"last_annual": "20251231", "name": "某公司",
                     "price": 38.0, "quote_source": "tencent",
                     "beta": 1.2, "beta_raw": 1.4, "beta_n": 104},
            "stmts": {"income": frame({"营业总收入": items.get("rev"),
                                       "净利润": items.get("np")}),
                      "balance": frame({"资产总计": items.get("ta"),
                                        "负债合计": items.get("tl")})}}


def _xc(cal, led):
    return dcf.cross_check(cal, led, AS_OF, _D)


# ── 一、两套数据源必须交叉核对 ───────────────────────────────

@needs_ashare
def test_单位换算正确():
    """ashare-dcf 以**亿元**计，账本以**元**计。换算错了会把每一项都
    报成天文数字的差异，或者反过来把真实差异淹掉。"""
    led = _led(revenue=97_900_000_000)          # 979 亿
    assert _xc(_cal(rev=979.0), led) == []


@needs_ashare
def test_对不上就报出来():
    """研报正文写一个数、估值模型用另一个数而无人知晓，是这套系统最不该
    出现的事 —— 与此前「权益乘数 1.77 对 1.94」同一类，只是跨了系统边界。"""
    led = _led(revenue=97_900_000_000)
    out = _xc(_cal(rev=950.0), led)
    # 报的是**实际匹配到的**科目名（营业总收入），不是候选列表里的第一个
    # —— 读者要能据此去报表里找到那一行。
    assert len(out) == 1 and "营业总收入" in out[0]
    assert "979" in out[0] and "950" in out[0] and "3.0%" in out[0]


@needs_ashare
def test_门槛设得紧():
    """两边取的是同一份公告数字，差异只会来自口径或单位，不会来自舍入。"""
    assert dcf.XCHECK_TOL <= D("0.01")
    led = _led(revenue=100 * 1e8)
    assert _xc(_cal(rev=100.9), led)      # 差 0.9%
    assert _xc(_cal(rev=100.2), led) == []  # 差 0.2%


@needs_ashare
def test_逐项核对而非只看一项():
    led = _led(revenue=100 * 1e8, net_profit=10 * 1e8,
               total_assets=200 * 1e8, total_liabilities=80 * 1e8)
    out = _xc(_cal(rev=100, np=8, ta=200, tl=70), led)
    assert len(out) == 2 and any("净利润" in x for x in out) \
        and any("负债合计" in x for x in out)


@needs_ashare
def test_账本缺该项时跳过而非误报():
    assert _xc(_cal(rev=979.0), FactLedger()) == []


@needs_ashare
def test_模型缺该项时跳过():
    led = _led(revenue=97_900_000_000)
    assert _xc(_cal(), led) == []


def test_没有年报期间就说不能核对():
    out = dcf.cross_check({"meta": {}, "stmts": {}}, _led(revenue=1), AS_OF, _D)
    assert out and "无法交叉核对" in out[0]


# ── 二、as_of：不报错的未来函数 ──────────────────────────────

def test_asof早于今天必须说出来():
    """ashare-dcf 取的是**当日**行情与最新报表。as_of 早于今天时它用的是
    未来信息 —— 这不是它的缺陷（它本就为「今天做一份模型」而设计），
    但两者接在一起时不说，就是一个不报错的未来函数。"""
    note = dcf._asof_note(date.today() - timedelta(days=30))
    assert note and "未来" not in note[0]        # 措辞不必吓人，但要说清
    assert "今日" in note[0] and "30 天" in note[0]


def test_asof就是今天则无需提示():
    assert dcf._asof_note(date.today()) == []


# ── 结果与落账本 ─────────────────────────────────────────────

def _res(**kw):
    base = dict(code="600584", workbook="/tmp/x.xlsx", price=D("38"),
                value_per_share=D("45.6"), wacc=D("0.0912"),
                terminal_roic=D("0.11"), reinvest_rate=D("0.42"),
                terminal_share=D("0.68"), verified=True)
    base.update(kw)
    return dcf.DcfResult(**base)


def test_空间由价值与现价算出():
    assert _res().upside == D("45.6") / D("38") - 1


def test_现价缺失时不给空间():
    assert _res(price=D(0)).upside is None


def test_摘要只用占位符():
    """与其它进提示词的摘要同一条约定 —— 提示词里出现裸数字、正文里却
    禁止引用，是自相矛盾的契约。"""
    from ir_agent.citation import audit_bare_numbers
    assert audit_bare_numbers(_res().summary()) == []


def test_验算未通过要写死在摘要里():
    txt = _res(verified=False).summary()
    assert "验算未通过" in txt and "不可采信" in txt


def test_估值结果落账本():
    fs = {f.key: f for f in dcf.to_facts(_res(), AS_OF)}
    assert fs["valuation.value_per_share"].value == D("45.6")
    assert fs["valuation.wacc"].unit == "ratio"
    assert all(f.period == AS_OF.isoformat() for f in fs.values())


def test_结果记为计算而非估计():
    """假设本身是判断，但由假设到结果这一步是确定性可复算的 ——
    工作簿就在那里，公式引擎跑过。"""
    assert all(f.method is Method.COMPUTED for f in dcf.to_facts(_res(), AS_OF))


# ── 找不到 ashare-dcf 时如实报告 ─────────────────────────────

def test_目录不存在时给出可操作的说明(tmp_path):
    with pytest.raises(dcf.DcfUnavailable) as e:
        dcf.run("600584", tmp_path, as_of=AS_OF,
                dirpath=Path(tmp_path) / "不存在")
    assert dcf.ENV_DIR in str(e.value)


def test_默认位置是同级目录():
    assert dcf.DEFAULT_DIR.name == "ashare-dcf"


@needs_ashare
def test_取值走ashare自己的换算而不另写一份():
    """`val()` 读原始值并除以 1e8 —— 换算逻辑只此一份。在这边照抄一遍，
    就是 ebit()、parent_first() 那类「一个概念两处实现」的老毛病。"""
    import inspect
    assert "D.val(" in inspect.getsource(dcf._stmt_value)
    assert "/ YI" not in inspect.getsource(dcf._stmt_value)


@needs_ashare
def test_科目名给候选而非写死():
    """不同公司列名略有出入（营业总收入 / 营业收入）。取不到就跳过，
    不误报成差异。"""
    names = dict((k, v) for k, _, v in
                 ((a, b, c) for a, b, c in dcf.XCHECK_ITEMS))
    assert "营业总收入" in names["revenue"] and "营业收入" in names["revenue"]


# ── 估值要先算，正文才谈得上「解释这个结果」 ─────────────────

def test_建模前必须有人批准():
    """**这是两阶段存在的理由。** 判断项（营收增速、毛利率）模型推不出来，
    只能靠分析得出 —— 也正因此必须有人看一眼再建模。第二阶段没有署名
    直接拒绝。"""
    import inspect

    from ir_agent import full
    src = inspect.getsource(full.run_valuation)
    assert "approved_by" in src
    assert "必须署名批准人" in src
    # 人改过的值要重新校验 —— 批准不等于豁免
    assert "assume.check(" in src


def test_完整研报只在建模之后产出():
    """建模之前就产出一份读起来像结论的完整研报，读者会把未经核实的假设
    当成定论。"""
    import inspect

    from ir_agent import full
    a = inspect.getsource(full.run_brief)
    v = inspect.getsource(full.run_valuation)
    assert "BRIEF" in a and "REPORT" not in a.replace("REPORT = ", "")
    assert "REPORT" in v and "dcf.run(" in v
    assert v.index("dcf.run(") < v.index("_write_full_report(")


def test_第八节要求解释结果而非重推假设():
    from ir_agent.analyst import _RULES
    assert "不是你定的" in _RULES and "不要给目标价" in _RULES
    assert "对账" in _RULES


def test_提示词点明强周期的适用边界():
    """拿周期中某一点的利润率外推永续，是 DCF 最常见的误用 ——
    而「强周期判给 PB–ROE」那道硬闸已随估值模块删除，这层提醒因此更要紧。"""
    from ir_agent.analyst import _RULES
    assert "强周期标的" in _RULES and "外推永续" in _RULES
    assert "DOL" in _RULES or "经营杠杆" in _RULES


def test_分解式非空时写入drivers块():
    """Task 10 接线：判断项按因子分解时（assumptions.decomps 非空），
    dcf.run 要把 assume.to_driver_block 的产出写进 cal["judg"]["drivers"]，
    ashare_build 才会建「1b_驱动分解」表。apply_to 必须先跑一遍——
    to_driver_block 的 base_level 依赖 cal["judg"]["gm_hist"]，那是
    calibrate 的原始输出，不受 apply_to 影响，但顺序仍要摆对，避免
    后来人误以为可以颠倒。"""
    import inspect

    src = inspect.getsource(dcf.run)
    assert "assume.apply_to(" in src
    assert "to_driver_block(" in src
    assert src.index("apply_to(") < src.index("to_driver_block(")
    assert "1b_驱动分解" in src
    assert 'cal["judg"]["drivers"]' in src.replace("'", '"')


def test_估值摘要进了提示词():
    from ir_agent.analyst import FactCatalog, build_prompt
    from ir_agent.ledger import FactLedger

    cat = FactCatalog(ledger=FactLedger(), as_of=AS_OF)
    txt = build_prompt(cat, code="600584", period="2025FY",
                       valuation_summary="DCF 估值（ashare-dcf 三表模型）")
    assert "ashare-dcf 三表模型" in txt and "估值结果" in txt


def test_估值摘要里的占位符必须能被替换机制识别():
    """**产出端写的占位符，消费端必须真能解析。**

    实测教训（600584 端到端）：`summary()` 此前写的是
    `[[valuation.wacc]]`——没有 `@period`，而 `citation.REF_RE` 硬性要求
    `[[key@period]]` 两段。于是这些占位符压根不被识别为占位符，一路原样
    漏进最终交付的 report.md：整个估值段落是六个未替换的占位符，真实数字
    （每股 28.92、WACC 9.11% 等）只出现在终端摘要里，读者拿到的研报没有
    一个估值数字。

    事实本身是入了账本的（`to_facts` 以 `as_of.isoformat()` 作期间），
    占位符也是照契约写的——只是两端对格式的约定不一致。这条测试钉住
    「摘要里的每一个占位符都能被 REF_RE 识别」，而不只是「摘要里有占位符」。
    """
    from datetime import date

    from ir_agent.citation import REF_RE

    as_of = date(2026, 9, 12)
    txt = _res(as_of=as_of).summary()
    refs = REF_RE.findall(txt)
    assert refs, "摘要里应当有可识别的占位符"
    keys = {k for _, k, _ in refs}
    assert keys >= {"valuation.value_per_share", "valuation.price",
                    "valuation.wacc", "valuation.terminal_roic",
                    "valuation.terminal_share"}, \
        f"估值关键项都要能被识别，实际只认出 {keys}"
    for _, _, per in refs:
        assert per == as_of.isoformat(), \
            f"期间必须与 to_facts 落账本时一致（{as_of.isoformat()}），得到 {per}"


def test_估值占位符的期间与落账本的期间一致():
    """回归锁：占位符里的 period 与 to_facts 写进账本的 period 必须同源。
    两边各写一次字面量，早晚会分叉——这条测试让分叉立刻可见。"""
    from datetime import date

    from ir_agent.citation import REF_RE
    from ir_agent import dcf

    as_of = date(2026, 9, 12)
    fact_periods = {f.period for f in dcf.to_facts(_res(as_of=as_of), as_of)}
    ref_periods = {per for _, _, per in REF_RE.findall(_res(as_of=as_of).summary())}
    assert fact_periods == ref_periods, \
        f"账本期间 {fact_periods} 与占位符期间 {ref_periods} 不一致"


class Test批准的数字必须真的进了工作簿:
    """**闸门查的都是「用得对不对」，没有一道查「结果对不对」。**

    这是与「占位符有没有被替换」同一类的漏网。人工批准判断项是整个流程
    的核心关卡，而批准之后没有任何一道检查回头确认：工作簿里真正参与
    估值的那几个数，就是人批准的那几个数。

    此前踩过的两次都是这条链断在中间（C2 / C-a）：同一因子同一年出现
    两行，清单显示其中一行、工作簿用另一行；重复行让 path() 按年份排序
    后整体错位一格，人批准第 2 年的 7%、工作簿拿到的是第 1 年那条重复行
    的 30%。两次都是**人自己读工作簿**才发现的。修法都是"不许出现重复
    行"——仍然是在查输入用得对不对，没有查输出对不对。

    `ashare_verify.read_inputs()` 本来就会把工作簿里逐年的营收增速/毛利率
    读回来（`growth`/`gm`，已含 Δ 偏移），拿它与 `apply_to` 写进
    `cal["judg"]` 的三档比一遍即可闭环，成本几乎为零。
    """

    def _judg(self):
        return {"growth": [0.10, 0.08], "gm": [0.15, 0.16],
                "growth_scen": {"bull": [0.14, 0.12], "base": [0.10, 0.08],
                                "bear": [0.05, 0.03]},
                "gm_scen": {"bull": [0.18, 0.19], "base": [0.15, 0.16],
                            "bear": [0.12, 0.11]}}

    def test_对得上时不报任何问题(self):
        from ir_agent import dcf
        inp = {"growth": [0.10, 0.08], "gm": [0.15, 0.16], "dg": 0.0, "dm": 0.0}
        assert dcf._readback_check(self._judg(), inp) == []

    def test_Δ偏移不算对不上(self):
        """`1_假设` 里的「营收增速 Δ」是人有意加的整体偏移，`read_inputs`
        把它并进了 growth。它不是错位，扣掉再比。"""
        from ir_agent import dcf
        inp = {"growth": [0.12, 0.10], "gm": [0.15, 0.16],
               "dg": 0.02, "dm": 0.0}
        assert dcf._readback_check(self._judg(), inp) == []

    def test_工作簿的数与批准的对不上要报出来(self):
        """整体错位一格正是踩过的那次：人批准第 1 年 10%、第 2 年 8%，
        工作簿拿到的是 8%、8%。"""
        from ir_agent import dcf
        inp = {"growth": [0.08, 0.08], "gm": [0.15, 0.16],
               "dg": 0.0, "dm": 0.0}
        problems = dcf._readback_check(self._judg(), inp)
        assert problems, "工作簿用的数与批准的不一致，必须报出来"
        assert "营收增速" in problems[0]
        assert "不可采信" in "".join(problems)

    def test_选中的是别的档要说清是哪一档(self):
        """工作簿的「选择情景」被改成乐观档时，数字对得上三档里的一档、
        只是不是中性档——这与"数字错了"是两回事，要分开说，否则人会以为
        建模坏了。"""
        from ir_agent import dcf
        inp = {"growth": [0.14, 0.12], "gm": [0.18, 0.19],
               "dg": 0.0, "dm": 0.0}
        problems = dcf._readback_check(self._judg(), inp)
        assert problems
        assert "乐观" in "".join(problems) or "bull" in "".join(problems)

    def test_年数对不上也要报(self):
        from ir_agent import dcf
        inp = {"growth": [0.10], "gm": [0.15, 0.16], "dg": 0.0, "dm": 0.0}
        assert dcf._readback_check(self._judg(), inp)


def test_回读校验真的接在建表之后():
    """光有 `_readback_check` 不算数——它必须真的被 `run()` 调用，而且要
    用建完表之后回读的 `inp`（不是建表前的 cal），否则查的还是输入。

    照本仓库既有做法（见 test_预研排在数据采集之前）用源码断言接线：
    这一步要真跑一遍得先有 ashare-dcf 与一个真实工作簿，代价太大。"""
    import inspect
    from ir_agent import dcf

    src = inspect.getsource(dcf.run)
    assert "_readback_check(" in src, "回读校验没有接进 run()"
    assert src.index("read_inputs") < src.index("_readback_check("), \
        "必须用建表之后回读的 inp"
    assert "assumptions is not None" in src


class Test锚定会覆盖人批准的首年假设:
    """**生益科技实跑当场抓到的盲区。**

    `1_假设` 的生效值 E 列公式是

        E11 = IF($B$30=1, B35/$B$81-1, CHOOSE($B$6,B11,C11,D11)+$B$26)
        E16 = IF($B$30=1, 1-B36/B35,  CHOOSE($B$6,B16,C16,D16)+$B$27)

    锚定开关 B30=1 时，**首年的营收增速与毛利率完全不用人批准的值**，
    改用已披露中期实际倒推。生益实测：人批准 22.31% / 27.77%，工作簿
    实际用 39.85% / 30.69%——而待核实清单、终端、研报全都按 22.31% 讲。

    `_readback_check` 上一版比的是 `read_inputs` 的 `growth`/`gm`，那是
    `path3()` 从 B/C/D **情景列**读的，情景列当然对得上（`apply_to` 刚
    写进去的）。**查了输入，没查生效值**——正是这道闸门本来要防的事，
    它自己犯了一遍。

    锚定本身是对的（用已披露的真实中报，比判断可靠）。要防的不是锚定，
    是"人批准的数被悄悄换掉而没人知道"。
    """

    def _judg(self):
        return {"growth": [0.2231, 0.1782], "gm": [0.2777, 0.2917],
                "growth_scen": {"bull": [0.45, 0.38], "base": [0.2231, 0.1782],
                                "bear": [0.03, 0.01]},
                "gm_scen": {"bull": [0.30, 0.33], "base": [0.2777, 0.2917],
                            "bear": [0.25, 0.24]}}

    def _inp(self, **kw):
        # 生益实测的真实数（ashare_verify.read_inputs 的 it 形状）：H1 营收 190.2561、
        # 成本 131.8622；按季度估算（2026-10-18 起）全年营收 426.5964、成本 293.2045
        # ——首年营收增速 = 年初至今同比 +50.05%，毛利率 31.27%；基期营收 284.3114
        it = {"sw": 1.0, "k": 2, "rev": 190.2561, "cogs": 131.8622,
              "fy_rev": 426.596379280956, "fy_cogs": 293.2044922458345}
        it.update(kw.pop("it", {}))
        d = {"growth": [0.2231, 0.1782], "gm": [0.2777, 0.2917],
             "dg": 0.0, "dm": 0.0, "it": it,
             "o": {"基期营业总收入": 284.3113845944}}
        d.update(kw)
        return d

    def test_锚定覆盖首年时必须报出来(self):
        from ir_agent import dcf
        problems = dcf._readback_check(self._judg(), self._inp())
        joined = "".join(problems)
        assert problems, "首年被锚定换掉却一声不吭，等于人工批准被架空"
        assert "锚定" in joined
        assert "22.31%" in joined and "50.05%" in joined, "两个数都要摆出来"
        assert "27.77%" in joined and "31.27%" in joined

    def test_锚定不是错_不该说不可采信(self):
        """锚定用的是已披露的中报实际，比判断可靠。要报的是「被换掉了」，
        不是「算错了」——把它说成不可采信会让人去关掉锚定，那才是坏事。"""
        from ir_agent import dcf
        joined = "".join(dcf._readback_check(self._judg(), self._inp()))
        assert "不可采信" not in joined

    def test_锚定开着也要照常查第二年起(self):
        """首年豁免，不等于整列豁免。第 2 年起仍然走 CHOOSE(情景)，
        对不上就是真的错位。"""
        from ir_agent import dcf
        bad = self._inp(growth=[0.2231, 0.9999])
        joined = "".join(dcf._readback_check(self._judg(), bad))
        assert "不可采信" in joined and "营收增速" in joined

    def test_锚定关闭时维持原样(self):
        from ir_agent import dcf
        off = self._inp(it={"sw": 0.0})
        assert dcf._readback_check(self._judg(), off) == []

    def test_没有锚定字段时维持原样(self):
        from ir_agent import dcf
        d = self._inp(); d["it"] = None
        assert dcf._readback_check(self._judg(), d) == []


class Test锚定说明要能单独取出来:
    """cd142a1 把「首年被锚定覆盖」塞在回读校验的 problems 里，只进得了终端。
    研报附录和第一阶段批准清单也要说同一句话，得有一个单独的出口；而反推
    公式只能有一份——三处各算一遍，早晚算出三个不同的数。"""

    _judg = Test锚定会覆盖人批准的首年假设._judg
    _inp = Test锚定会覆盖人批准的首年假设._inp

    def test_反推公式是一个纯函数(self):
        g, m = dcf.anchored_year1(fy_rev=426.596379280956, fy_cogs=293.2044922458345,
                                  base_rev=284.3113845944)
        assert g == pytest.approx(0.5005, abs=5e-4)
        assert m == pytest.approx(0.3127, abs=5e-4)

    def test_锚定说明单独可取且与回读校验说的是同一句(self):
        note = dcf._anchor_note(self._judg(), self._inp())
        assert "50.05%" in note and "31.27%" in note
        assert "22.31%" in note and "27.77%" in note
        assert "不可采信" not in note
        assert note in dcf._readback_check(self._judg(), self._inp())

    def test_没锚定时没有说明(self):
        assert dcf._anchor_note(self._judg(), self._inp(it={"sw": 0.0})) == ""

    def test_建模结果带出锚定说明(self):
        import inspect
        assert dcf.DcfResult.__dataclass_fields__["anchor_note"].default == ""
        assert "anchor_note=" in inspect.getsource(dcf.run)


class Test研报对WACC的归因必须指向真正参与计算的β:
    """**生益科技研报原话**：「模型所用WACC为9.84%，其中权益成本部分依赖的β
    （对沪深300、周频两年）为1.84x」。而 9.84% 是用 1.3717（五年对上证综指 +
    Blume）算的；用 1.84 算，WACC 约 12.6%。交付物说了假话。

    根子：账本里只有 ir-agent 自己回归的 `beta`，模型实际用的 β 从未落账本，
    研报想谈 WACC 的 β 只引得到那一个。"""

    def test_模型所用β落账本(self):
        res = _res(as_of=AS_OF)
        res.beta = D("1.3717")
        keys = {f.key: f for f in dcf.to_facts(res, AS_OF)}
        assert keys["valuation.beta"].value == D("1.3717")
        assert keys["valuation.beta"].period == AS_OF.isoformat()

    def test_摘要里给出模型所用β的占位符(self):
        res = _res(as_of=AS_OF)
        res.beta = D("1.3717")
        assert f"[[valuation.beta@{AS_OF.isoformat()}]]" in res.summary()

    def test_另一个β的标签写明未用于WACC(self):
        from ir_agent.analyst import LABELS
        assert "未用于 WACC" in LABELS["beta"]
        assert "WACC" in LABELS["valuation.beta"]

    def test_建模时从工作簿读出所用β(self):
        import inspect
        assert 'inp["dcf"]["Beta（5年"]' in inspect.getsource(dcf.run)


# ── 量级对账：每股价值与现价不能差一个数量级 ──────────────────────
#
# **实测事故（生益科技 600183，第二阶段首次跑通）。** 取数降级到腾讯通道，
# 那条分支把换手率（1.47）当成总股本读走，总股本从 24.29 亿股变成
# 0.000147 亿股，每股内在价值算出 25,077,036 元 —— 现价 129.96。
#
# 根因已在 ashare-dcf 的 `_check_quote` 修掉（现价 × 股本 要对得上同一条
# 响应里的总市值）。这里加的是**第二道网，且不依赖病因**: 任何原因导致
# 每股价值与现价差一个数量级，都该先怀疑数据而不是先怀疑市场。
#
# 为什么两道网都要有: 上游那条对的是取数口径，它拦不住建模层的单位错位；
# 这条对的是最终输出，它不知道错在哪但知道结果不可信。一道网只在它所在
# 的那一层有效。
class TestScaleCheck:
    def test_每股价值是现价的十九万倍要拦下(self):
        w = dcf._scale_warnings(D("25077036.37"), D("129.96"))
        assert w and "不可采信" in w[0]

    def test_每股价值只有现价的万分之一也要拦下(self):
        """反面同样是量级错位 —— 单位少乘一次和多乘一次同样常见。"""
        assert dcf._scale_warnings(D("0.013"), D("129.96"))

    def test_正常的上行空间不该被拦(self):
        assert dcf._scale_warnings(D("151.70"), D("129.96")) == []

    def test_翻三倍的深度低估不该被拦(self):
        """DCF 算出三倍空间是罕见但真实的结论，不是数据错 —— 这条钉住
        阈值别收得太紧，否则这道网会在真结论上响，然后被人忽略。"""
        assert dcf._scale_warnings(D("114"), D("38")) == []

    def test_现价缺失时不判(self):
        assert dcf._scale_warnings(D("151.7"), D("0")) == []

    def test_量级告警排在其它健全性告警前面(self):
        """数要是根本不对，后面那些告警都是在给一个错的模型体检 ——
        与 `_readback_check` 同一条道理。"""
        res = _res(price=D("129.96"), value_per_share=D("25077036.37"),
                   warnings=["终值占 EV 94%：估值几乎全部来自终值假设"])
        got = dcf._with_scale_check(res)
        assert "不可采信" in got.warnings[0]
        assert "终值占 EV" in got.warnings[1]


def test_量级对账真的接在run里():
    """**这条钉的是「接上了」，不是「写出来了」。** 本轮另一处踩过同样的
    坑: 传输重试的计数器写好了却没人读，注释里却声称产物会报告它。一个
    没有调用点的检查函数，和没写这个函数没有区别。

    这里查的是调用点存不存在（函数名比 except 子句的字面量稳定得多），
    真正的行为由上面 TestScaleCheck 那组测试钉住。
    """
    import inspect

    assert "_with_scale_check(res)" in inspect.getsource(dcf.run)


# ── 终值占 EV 接近 100% = DCF 已退化成一个终值公式 ────────────────
#
# **生益科技 600183 实跑（2026-10-06）。** 判断项给出五年 8 倍营收，资本开支
# 占营收 15–19%，于是明确期 FCFF 现值只剩 2.26 亿，终值占 EV **100.0%** ——
# 而终端照样印出「每股内在价值 183.27 元 vs 现价 129.96 元，空间 +41.0%」，
# 研报也照样产出了。
#
# 既有的健全性检查报了这一条（>90% 那档），但它和「终值占 78%」同级，只是
# 清单里的一条 ⚠。终值占到 100% 不是「高度依赖终值」，是**明确期不贡献任何
# 价值**: 这时 DCF 已经退化成 `FCFF终/(WACC−g)` 一个公式，五年三表预测对
# 结果毫无影响，而那五年正是全部判断项的用处所在。
#
# 与 `_scale_warnings` 同一处理（排在最前、措辞「不可采信」），理由也一样:
# 数要是根本不对，后面那些告警都是在给一个错的模型体检。
class TestTerminalShareBlocks:
    def test_终值占满百判为不可采信(self):
        w = dcf._terminal_warnings(D("1.0"))
        assert w and "不可采信" in w[0]

    def test_九成八也算退化(self):
        assert dcf._terminal_warnings(D("0.985"))

    def test_八成四不触发(self):
        """茅台 87.8%、生益 b6 轮 83.9% 都是真实可用的结果 —— 高但不退化。
        这条钉住阈值别收得太紧: 一个常响的「不可采信」会被学会忽略。"""
        assert dcf._terminal_warnings(D("0.878")) == []
        assert dcf._terminal_warnings(D("0.839")) == []

    def test_接进了run(self):
        import inspect
        src = inspect.getsource(dcf.run)
        assert "_terminal_warnings(" in src
        assert src.index("_terminal_warnings(") < \
            src.index('res.warnings.insert(0, f"公式引擎验算失败')


# ── 蒙特卡洛：dcf.py 声称选用 ashare-dcf 的理由之一，却从未接上 ──────────
#
# ashare-dcf 命令行在验算通过后、带 --mc 时追加「8_蒙特卡洛」页；`dcf.run` 照抄了
# 前三步（建表、验算、健全性），漏了这一段 —— git 历史里 ir_agent 从未引用过
# ashare_mc。生益科技 600183 实测：工作簿 8 张表，没有蒙特卡洛页。

def _mc_cal():
    return {"judg": {"g_stats": (0.1082, 0.2121, -0.1115, 0.3945),
                     "gm_stats": (0.2331, 0.0292, 0.1919, 0.3069)},
            "_rho": 0.385, "_rho_n": 30,
            "gen": {"tax": 0.104}, "dcf": {"beta": 1.3735}}


def _mc_out(**kw):
    """ashare_mc.run 的返回形状（键名照抄其源码）。"""
    base = dict(n=50000, seed=42, base=56.56, mean=61.92, p50=55.54, sd=29.89,
                p5=28.71, p95=115.59, p_above=0.0445, clipped=0.044, se=0.134,
                degraded=0.0, neg_ebit=0.0, price=118.77, src="首次生成，使用默认参数",
                tornado=[{"name": "Beta → WACC", "low": 40.0, "high": 76.2, "swing": 36.2},
                         {"name": "营收增速（路径平移）", "low": 41.0, "high": 75.2, "swing": 34.2},
                         {"name": "永续增长率 g", "low": 50.0, "high": 62.2, "swing": 12.2},
                         {"name": "毛利率（路径平移）", "low": 53.0, "high": 59.0, "swing": 6.0}])
    base.update(kw)
    return base


class _FakeMC:
    def __init__(self, out=None, exc=None):
        self.out, self.exc, self.calls = out, exc, []

    def run(self, path, hist=None, n=None, seed=None, quiet=False):
        self.calls.append(dict(path=path, hist=hist, quiet=quiet))
        if self.exc:
            raise self.exc
        return self.out


def _mc_res(**kw):
    return _res(as_of=AS_OF, mc=dcf._monte_carlo(_FakeMC(_mc_out()), "/tmp/x.xlsx",
                                                 _mc_cal())[0], **kw)


class Test蒙特卡洛:
    def test_历史波动的取法与命令行一致(self):
        """命令行 `ashare_dcf.py --mc` 怎么从 cal 取，这里就怎么取 —— 不另发明一套。"""
        assert dcf._mc_hist(_mc_cal()) == dict(
            sd_growth=0.2121, sd_margin=0.0292, rho=0.385, rho_n=30,
            tax_mean=0.104, beta_mean=1.3735)

    def test_用公司历史波动静默运行(self):
        mc = _FakeMC(_mc_out())
        dcf._monte_carlo(mc, "/tmp/x.xlsx", _mc_cal())
        assert mc.calls[0]["hist"] == dcf._mc_hist(_mc_cal())
        assert mc.calls[0]["quiet"] is True

    def test_取出分布与龙卷风前三(self):
        got, warns = dcf._monte_carlo(_FakeMC(_mc_out()), "/tmp/x.xlsx", _mc_cal())
        assert warns == []
        assert got["p5"] == D("28.71") and got["p50"] == D("55.54")
        assert got["p95"] == D("115.59") and got["mean"] == D("61.92")
        assert got["p_above"] == D("0.0445") and got["n"] == D(50000)
        # 摆幅取一半，与命令行打印的「±」一致；只要前三
        assert got["tornado"] == [("Beta → WACC", D("18.1")),
                                  ("营收增速（路径平移）", D("17.1")),
                                  ("永续增长率 g", D("6.1"))]

    def test_失败不拖垮估值但要说出来(self):
        """蒙特卡洛是附加层：引擎自检不一致或依赖缺失时，点估值照常交付，
        但不能静默 —— 第 8 节要知道没有模拟结果，不能编一个。"""
        got, warns = dcf._monte_carlo(
            _FakeMC(exc=AssertionError("蒙特卡洛引擎与确定性模型不一致")),
            "/tmp/x.xlsx", _mc_cal())
        assert got == {}
        assert warns and "蒙特卡洛未生成" in warns[0] and "AssertionError" in warns[0]

    def test_缺历史波动就不跑而不是退回默认值(self):
        """默认 σ（增速 15%、毛利率 3%）不是这家公司的分布 —— 用它跑出来的
        P5/P95 读起来像结论，实际是模板参数。"""
        mc = _FakeMC(_mc_out())
        got, warns = dcf._monte_carlo(mc, "/tmp/x.xlsx", {"judg": {}})
        assert got == {} and mc.calls == []
        assert "蒙特卡洛未生成" in warns[0]

    def test_摘要里的模拟数字全是可识别的占位符(self):
        from ir_agent.citation import REF_RE, audit_bare_numbers
        txt = _mc_res().summary()
        assert "蒙特卡洛" in txt
        assert audit_bare_numbers(txt) == []
        keys = {k for _, k, _ in REF_RE.findall(txt)}
        assert {"valuation.mc_p5", "valuation.mc_p50", "valuation.mc_p95",
                "valuation.mc_p_above"} <= keys

    def test_摘要里每个占位符都落了账本(self):
        """产出端写什么，账本就得有什么 —— 600584 那次整段估值漏成未替换占位符
        的同一类问题，这里对蒙特卡洛新增的每个键再钉一次。"""
        from ir_agent.citation import REF_RE
        res = _mc_res()
        refs = {(k, p) for _, k, p in REF_RE.findall(res.summary())}
        facts = {(f.key, f.period) for f in dcf.to_facts(res, AS_OF)}
        assert refs <= facts, f"摘要引用了账本里没有的事实：{refs - facts}"

    def test_模拟结果落账本(self):
        fs = {f.key: f for f in dcf.to_facts(_mc_res(), AS_OF)}
        assert fs["valuation.mc_p5"].value == D("28.71")
        assert fs["valuation.mc_p5"].unit == "元"
        assert fs["valuation.mc_p_above"].unit == "ratio"
        assert fs["valuation.mc_swing_1"].value == D("18.1")

    def test_路径数渲染后自带完整量词(self):
        """生益科技试跑：摘要写「[[mc_paths]]路径」，单位是「条」，模型照抄时在
        占位符后补了空格，研报成了「50000 条 路径」。单位整体定成「条路径」、
        摘要里占位符后不再跟名词，就没有可补空格的位置。"""
        from ir_agent.citation import format_value
        fs = {f.key: f for f in dcf.to_facts(_mc_res(), AS_OF)}
        assert format_value(fs["valuation.mc_paths"]) == "50000 条路径"
        assert "]]路径" not in _mc_res().summary()

    def test_没有模拟结果时不落账本也不进摘要(self):
        assert not [f for f in dcf.to_facts(_res(), AS_OF)
                    if f.key.startswith("valuation.mc_")]
        assert "valuation.mc_" not in _res(as_of=AS_OF).summary()

    def test_接在验算通过之后(self):
        """命令行在验算未通过时直接返回、不跑蒙特卡洛；这里同一顺序。"""
        import inspect
        src = inspect.getsource(dcf.run)
        assert "import ashare_mc" in src and "_monte_carlo(" in src
        assert src.index("V.verify(") < src.index("_monte_carlo(")
        assert "if not fails" in src

    @needs_ashare
    def test_真实引擎把蒙特卡洛页写进工作簿(self, tmp_path):
        """用 ashare-dcf 自带的生益科技校准固件建一份真工作簿，跑真引擎。"""
        import copy
        import pickle

        import ashare_build as B
        import ashare_mc as MC
        from openpyxl import load_workbook

        fix = dcf.DEFAULT_DIR / "tests" / "fixtures" / "cal_600183.pkl"
        with open(fix, "rb") as f:
            cal = pickle.load(f)
        wb = tmp_path / "w.xlsx"
        B.build(copy.deepcopy(cal), str(wb))

        got, warns = dcf._monte_carlo(MC, str(wb), cal)
        assert warns == []
        book = load_workbook(wb)
        assert MC.SHEET in book.sheetnames
        assert got["p5"] < got["p50"] < got["p95"]
        # 参数区用的是公司历史波动，不是模板默认值
        sd_g = MC.read_params(book[MC.SHEET])["sd_g"]
        assert sd_g == MC.defaults(dcf._mc_hist(cal))["sd_g"]
        assert sd_g != MC.defaults({})["sd_g"]


def test_第八节要求解读蒙特卡洛分布():
    from ir_agent.analyst import FactCatalog, _RULES, build_prompt
    from ir_agent.ledger import FactLedger

    assert "蒙特卡洛" in _RULES and "高于现价的概率" in _RULES
    # 模拟的波动取自历史，不是人批准的三档 —— 两者回答的问题不同，必须点明
    assert "历史波动" in _RULES
    cat = FactCatalog(ledger=FactLedger(), as_of=AS_OF)
    txt = build_prompt(cat, code="600584", period="2025FY",
                       valuation_summary="DCF 估值（ashare-dcf 三表模型）")
    assert "蒙特卡洛" in txt.split("**估值结果**")[1]
