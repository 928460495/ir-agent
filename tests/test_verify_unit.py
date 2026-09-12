# -*- coding: utf-8 -*-
"""ashare_verify 的纯函数单元测试：不依赖真实工作簿，只测 _cell_value / _eval_arith
两条边角逻辑。真实 140 格独立复算回归见 tests/test_financing.py（脚本方式跑）。

对应终审两条发现：
1) 循环引用守卫的 ValueError 被 _eval_arith 的 except (TypeError, ValueError) 吞掉，
   诊断信息（"引用层数超过 5 —— 可能存在循环引用"）丢失，下游变成一个无信息量的
   float() 类型错误。
2) eval 字符白名单允许相邻的 "*" 拼成 "**"（Python 乘方），Excel 自己的乘方符号是
   "^"（已被白名单挡掉），但被人手改的工作簿可以写出 9**9**9 这种指数塔拖死进程。
"""
import os, sys
import pytest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import Workbook
import ashare_verify as V


def _mk(cells):
    """建一个只写了 A 列若干行公式/常量的内存工作簿。"""
    wb = Workbook(); ws = wb.active
    for addr, v in cells.items():
        ws[addr] = v
    return wb, ws


# ---------- 1) 循环引用诊断信息不应被吞掉 ----------

def test_circular_reference_via_arith_raises_with_diagnostic():
    """A1..A7 首尾相连、且每一环都是『四则运算』公式（=A2+0 而非纯引用 =A2），
    这样才会经过 _eval_arith。修复前：_eval_arith 的 except (TypeError, ValueError)
    吞掉诊断，向上冒泡出来的是字符串 '=A2+0'（_cell_value 直接 return v），
    下游 float() 会报一个看不出「循环引用」四个字的 TypeError/ValueError。
    修复后应直接抛出带「循环引用」字样的 ValueError。"""
    wb, ws = _mk({
        "A1": "=A2+0", "A2": "=A3+0", "A3": "=A4+0", "A4": "=A5+0",
        "A5": "=A6+0", "A6": "=A7+0", "A7": "=A1+0",
    })
    with pytest.raises(ValueError, match="循环引用"):
        V._cell_value(wb, ws, "A1")


def test_circular_reference_via_pure_ref_chain_still_raises():
    """回归锁定：纯引用链（=B2 这种整格引用，不经过 _eval_arith）本来就没有被吞掉过，
    修复不能把这条路弄坏。"""
    wb, ws = _mk({
        "A1": "=A2", "A2": "=A3", "A3": "=A4", "A4": "=A5",
        "A5": "=A6", "A6": "=A7", "A7": "=A1",
    })
    with pytest.raises(ValueError, match="循环引用"):
        V._cell_value(wb, ws, "A1")


def test_eval_arith_parse_failure_still_falls_back_to_none():
    """回归锁定：_eval_arith 的 except 分支不能被整个删掉——解析不了的算式（如函数调用）
    仍必须走 return None 的老路，让 _cell_value 把原始公式串原样吐回去，
    由下游 float() 响亮崩溃，而不是让这里静默出错值或索性崩在这一层。"""
    wb, ws = _mk({"A1": "=SUM(B1:B2)", "B1": 1, "B2": 2})
    v = V._cell_value(wb, ws, "A1")
    assert v == "=SUM(B1:B2)"


def test_eval_arith_ordinary_formula_unaffected():
    """回归锁定：普通四则运算公式（含跨单元格引用）不受这次修复影响。"""
    wb, ws = _mk({"A1": "=B1+C1", "B1": 3, "C1": 4})
    assert V._cell_value(wb, ws, "A1") == 7.0


# ---------- 2) eval 白名单必须显式拒绝 "**" ----------

def test_eval_arith_rejects_double_star():
    """白名单按字符放行 '*'，相邻两个就拼成 Python 的乘方 '**'。Excel 自己的乘方符号
    是 '^'（已被白名单挡掉），只有手改的工作簿才能写出这种式子；9**9**9 这种指数塔
    会算出一个几亿位的大整数，拖死校验进程（已实测挂起超时）。这里只用小指数验证
    守卫生效，不实际触发爆算。"""
    assert V._eval_arith(Workbook(), Workbook().active, "2**3", 0) is None


def test_eval_arith_single_star_still_works():
    """回归锁定：单个 '*'（乘法）不能被这次的显式拒绝误伤。"""
    wb = Workbook()
    assert V._eval_arith(wb, wb.active, "2*3", 0) == 6.0
