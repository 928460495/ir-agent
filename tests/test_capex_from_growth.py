# -*- coding: utf-8 -*-
"""资本开支由营收增速推出。

**为什么不能三档同值。** 原实现三档都用「近三年 capex/营收均值线性收敛到
折旧率」，于是 Bull 档营收增速 45% 也配同样的资本开支——等于假设增长不
需要投入。只调有利项是情景分析最常见的作弊方式。

公式：当年资本开支 = Δ营收 ÷ 目标固定资产周转率 + 维持性投资（≈折旧）。
"""
import copy, os, pickle, sys
import pytest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook
import ashare_calibrate as C
import ashare_build as B

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "cal_600183.pkl")


def test_增速越高资本开支越高():
    """三档自动一致：Bull 增速高 → 需要更多产能 → 资本开支更高。"""
    out = C._capex_from_growth(
        growth3={"bull": [0.45, 0.30], "base": [0.20, 0.15], "bear": [0.05, 0.03]},
        fa_turnover=2.0, maint_rate=0.05)
    assert out["bull"][0] > out["base"][0] > out["bear"][0]


def test_公式逐位可复算():
    """Δ营收 100×0.20 = 20；20 ÷ 2.0 = 10 增量投资；维持性 100×0.05 = 5；合计 15。
    以营收占比表示：15 / (100×1.20) = 0.125"""
    out = C._capex_from_growth(
        growth3={"base": [0.20]}, fa_turnover=2.0, maint_rate=0.05)
    assert abs(out["base"][0] - 0.125) < 1e-9, out["base"][0]


def test_零增长时只剩维持性投资():
    out = C._capex_from_growth(
        growth3={"base": [0.0]}, fa_turnover=2.0, maint_rate=0.05)
    assert abs(out["base"][0] - 0.05) < 1e-9


def test_负增长不产生负资本开支():
    """**收缩不等于卖厂房。** 营收下滑时增量投资归零，但维持性投资还在。"""
    out = C._capex_from_growth(
        growth3={"bear": [-0.10]}, fa_turnover=2.0, maint_rate=0.05)
    assert out["bear"][0] >= 0.05, out["bear"][0]


def test_连续多年负增长时资本开支比例会回升():
    """**这不是 bug。** I5（final-review-plan2a 终审）之后，公式化简为
    capex_ratio_t = (max(g_t,0)/fa_turnover + maint_rate)/(1+g_t)，不再有
    rev_prev/rev_now 的递推——旧版 docstring 声称这条测试能抓"递推顺序
    错误"，终审实跑两个突变体（删掉 rev_prev=rev_now；base_rev 写死 1.0）
    验证了旧实现里这套递推本来就不影响输出，声称落空。现在这条测试改为
    钉住化简后公式本身的形状：连续负增长时 g<0 → 分子恒为 maint_rate，
    只有分母 (1+g) 在缩小，g 越负、比例越高——单调走高是公式的直接推论，
    不是"上一年营收当基数"这类会被递推顺序搞错的东西。
    """
    out = C._capex_from_growth(
        growth3={"bear": [-0.05, -0.10, -0.15]}, fa_turnover=2.0, maint_rate=0.05)
    bear = out["bear"]
    # 逐位核算：maint_rate/(1+g) = 0.05/0.95, 0.05/0.90, 0.05/0.85
    assert bear == pytest.approx([0.0526, 0.0556, 0.0588], abs=1e-4)
    assert bear[0] < bear[1] < bear[2], bear


def test_各年比例只依赖当年增速与更早年份无关():
    """I5：capex_ratio_t 代数上只是 g_t 的函数——分子分母里的 rev_prev 恒等
    约掉，因此不管上一年增速是什么，只要当年 g 相同，比例就必须相同。删掉
    `base_rev` 参数之后，原 `test_逐年以上一年营收为基数` 名字里"基数"这个
    概念已经不存在了，改写成直接钉住这份代数化简，而不是继续用一个已经
    不成立的理由（"用上一年营收当基数"）去解释一个恰好也算得对的数字。
    """
    seq_a = C._capex_from_growth(growth3={"base": [0.20, 0.20]}, fa_turnover=2.0, maint_rate=0.05)
    seq_b = C._capex_from_growth(growth3={"base": [-0.30, 0.20]}, fa_turnover=2.0, maint_rate=0.05)
    assert seq_a["base"][1] == pytest.approx(seq_b["base"][1], abs=1e-9), (
        "第2年比例只应看第2年自己的增速 0.20，不该被第1年是 0.20 还是 -0.30 影响")
    # 第2年：g=0.20 → (0.20/2.0+0.05)/1.20 = 0.125，与基期无关
    assert abs(seq_a["base"][1] - 0.125) < 1e-9, seq_a["base"][1]


# ── calibrate 交付 fa_turn / maint_rate，供 build 使用 ──────────────────

def test_calibrate交付fa_turn和maint_rate(cal_dict):
    ju = cal_dict["judg"]
    assert "fa_turn" in ju and "maint_rate" in ju
    assert 0.3 <= ju["fa_turn"] <= 10.0
    assert ju["maint_rate"] >= 0.0
    # 旧的 judg["capex"] 单序列已删除（见 task-3-report.md Important 1）：
    # _scen 的 fallback 只作用于 growth_scen，capex 从未被传给它，全仓库
    # 无消费方，留着只会让下一个维护者去找一个不存在的用途。
    assert "capex" not in ju


def test_夹具与当前calibrate代码同步(cal_dict, monkeypatch):
    """夹具是某次 calibrate() 运行的离线快照（pickle），不会跟着代码改动
    自动更新。Task 3 审查时发现它已经滞后两三个提交——季节性对冲
    （cfabb70）与 gm_std_annual 修正（8e96197）落地时都没人重新生成过，
    直到本任务顺手 regenerate 才暴露（task-3-report.md Important 2）。
    没有测试盯着这件事，夹具就会悄悄过期而没人发现。

    M3（final-review-plan2a 终审）：原实现硬编码「fa_turn/maint_rate 必须在、
    capex/_hedge_note 必须在/不在」这三四条断言，对未来任何其它字段的增删
    都无感——加一个新字段、删一个旧字段，只要没人手改这条测试，它照样全绿。
    改成不预设任何具体字段名：用夹具里已经缓存的 stmts（三张报表）+ meta
    （行情/beta/无风险利率）离线重放一遍*当前代码*的 calibrate()——不联网、
    不需要真实调用东财/新浪，只把网络取数函数换成读夹具里存好的值——把
    重放结果的 judg/gen 键集合与夹具里存的键集合逐一比对。任何一侧多一个键
    或少一个键，都说明代码已经改了夹具没跟着变，或者反过来。
    """
    meta = cal_dict["meta"]
    monkeypatch.setattr(C.D, "fetch_statements", lambda code: cal_dict["stmts"])
    monkeypatch.setattr(C.D, "fetch_quote", lambda code: dict(
        name=meta["name"], price=meta["price"], shares=meta["shares"], source=meta["quote_source"]))
    monkeypatch.setattr(C.D, "fetch_riskfree", lambda: (meta["rf"], meta["rf_date"]))
    monkeypatch.setattr(C, "beta_blume", lambda code, years=5: (meta["beta"], meta["beta_raw"], meta["beta_n"]))
    # 东财折旧摊销明细源置空——calibrate 自带"东财不可用则降级到 OCF 倒推"的
    # 分支，两条分支产出的 judg/gen 键集合完全相同，只是数值精度不同；这里
    # 只关心键集合，降级分支足够、也不需要联网。
    monkeypatch.setattr(C.D, "fetch_eastmoney", lambda code, n=24: {"_errors": []})
    fresh = C.calibrate(meta["code"], meta["nyears"], stmts=cal_dict["stmts"])
    stale_msg = ("夹具 tests/fixtures/cal_600183.pkl 已落后于当前 "
                 "ashare_calibrate.calibrate()，请重新生成：python3 tests/make_fixture.py 600183")
    ju_missing = set(fresh["judg"]) - set(cal_dict["judg"])
    ju_extra = set(cal_dict["judg"]) - set(fresh["judg"])
    assert not ju_missing and not ju_extra, (
        f"judg 键集合不一致：当前代码多出 {ju_missing}，夹具多出 {ju_extra}——{stale_msg}")
    ge_missing = set(fresh["gen"]) - set(cal_dict["gen"])
    ge_extra = set(cal_dict["gen"]) - set(fresh["gen"])
    assert not ge_missing and not ge_extra, (
        f"gen 键集合不一致：当前代码多出 {ge_missing}，夹具多出 {ge_extra}——{stale_msg}")
    it_fresh, it_fix = fresh.get("interim") or {}, cal_dict.get("interim") or {}
    assert set(it_fresh) == set(it_fix), (
        f"interim 键集合不一致：当前代码 {sorted(set(it_fresh) - set(it_fix))}，"
        f"夹具 {sorted(set(it_fix) - set(it_fresh))}——{stale_msg}")


@pytest.fixture(scope="module")
def cal_dict():
    if not os.path.exists(FIX):
        pytest.skip("缺离线夹具：python3 tests/make_fixture.py 600183")
    with open(FIX, "rb") as f:
        return pickle.load(f)


def _rowmap(ws):
    m = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip():
            m.setdefault(v.strip(), r)
    return m


def _with_growth_scen(cal, bull, base, bear):
    """注入三档营收增速判断项（ir-agent 的 assume.py 产出形状）。"""
    c = copy.deepcopy(cal)
    c["judg"]["growth_scen"] = {"bull": bull, "base": base, "bear": bear}
    return c


def _anchored_g1(cal):
    """首个预测年按季度估算后的生效增速（三档相同）。"""
    return cal["interim"]["fy_rev"] / cal["rev"][cal["periods"]["annual"][-1]] - 1.0


def test_工作簿里三档资本开支真的拉开(cal_dict, tmp_path):
    """本任务要断言的是工作簿里真正生效的那几格，不是中间变量——
    上一个任务的教训：只断言留痕字典，测不出函数是否真的把值写回。

    夹具（600183）有已披露中期，锚定默认开启：首个预测年的增速已不是这里注入
    的 0.45/0.20/0.05，而是季度估算的生效增速（三档相同）——第 2 年起三档才按
    注入值拉开。逐位复算必须先做同一份首年覆盖再喂给 _capex_from_growth。
    """
    n = len(cal_dict["periods"]["proj"])
    bull = [0.45] * n; base = [0.20] * n; bear = [0.05] * n
    cal = _with_growth_scen(cal_dict, bull, base, bear)
    assert cal["interim"] and cal["interim"].get("fy_rev"), "前提：夹具确有已披露中期"
    p = tmp_path / "w.xlsx"
    B.build(cal, str(p))
    ws = load_workbook(p)["1_假设"]
    m = _rowmap(ws)
    proj = cal["periods"]["proj"]
    g1 = _anchored_g1(cal)
    growth3 = {"bull": [g1] + bull[1:], "base": [g1] + base[1:], "bear": [g1] + bear[1:]}
    expect = C._capex_from_growth(
        growth3=growth3, fa_turnover=cal["judg"]["fa_turn"],
        maint_rate=cal["judg"]["maint_rate"])
    for i, y in enumerate(proj):
        r = m[f"{y} 资本开支 / 营收"]
        got = [ws.cell(r, c).value for c in (2, 3, 4)]
        assert got == pytest.approx([expect[k][i] for k in ("bull", "base", "bear")], abs=1e-4), y
    r1 = m[f"{proj[1]} 资本开支 / 营收"]
    bull1, base1, bear1 = (ws.cell(r1, c).value for c in (2, 3, 4))
    assert bull1 > base1 > bear1, "第 2 年起三档按注入的增速拉开"


def test_首年资本开支用季度估算的生效增速_三档相同(cal_dict, tmp_path):
    """C2（final-review-plan2a 终审）的教训仍然成立：锚定开启时首年的增速判断项
    被完全覆盖，资本开支必须按真正生效的增速算，不能用那个不生效的判断项。
    2026-10-18 起首年改为季度估算，三档相同——首年资本开支也三档相同。"""
    cal = copy.deepcopy(cal_dict)
    p = tmp_path / "w.xlsx"
    B.build(cal, str(p))
    ws = load_workbook(p)["1_假设"]
    m = _rowmap(ws)
    r0 = m[f"{cal['periods']['proj'][0]} 资本开支 / 营收"]
    got = [ws.cell(r0, c).value for c in (2, 3, 4)]
    g1 = _anchored_g1(cal)
    expect = C._capex_from_growth(
        growth3={k: [g1] for k in ("bull", "base", "bear")},
        fa_turnover=cal["judg"]["fa_turn"], maint_rate=cal["judg"]["maint_rate"])
    assert got == pytest.approx([expect["base"][0]] * 3, abs=1e-4), got
    judged = cal["judg"]["growth"][0]
    assert abs(g1 - judged) > 1e-3, "前提：生效增速与被覆盖的判断项确实不同，否则测不出区别"
