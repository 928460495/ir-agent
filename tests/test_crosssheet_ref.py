"""ashare_verify 的跨表引用解析。

**这是验证核心。** read_inputs 是「独立复算」的读取器 —— 它读错了，
整个「公式引擎实跑验算」的担保就失效了。因此改动虽小，测试要密。
"""

from __future__ import annotations

import sys

import pytest
from openpyxl import Workbook

from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "ashare-dcf"))


def _wb(tmp_path, value_cell, ref):
    """造一个两表工作簿：assumptions 表引用 drivers 表。"""
    wb = Workbook()
    a = wb.active
    a.title = "1_假设"
    a["A1"] = "测试项"
    a["B1"] = ref
    d = wb.create_sheet("1b_驱动分解")
    d[value_cell] = 0.1234
    p = tmp_path / "x.xlsx"
    wb.save(p)
    return p


def test_解析跨表引用():
    from ashare_verify import _resolve_ref
    assert _resolve_ref("=1b_驱动分解!B5") == ("1b_驱动分解", "B5")
    assert _resolve_ref("='1b_驱动分解'!$B$5") == ("1b_驱动分解", "B5")


def test_本表引用仍然认得():
    """原有能力不能丢。"""
    from ashare_verify import _resolve_ref
    assert _resolve_ref("=B48") == (None, "B48")
    assert _resolve_ref("=$B$48") == (None, "B48")


def test_非引用返回空():
    from ashare_verify import _resolve_ref
    assert _resolve_ref("=B1+C1") is None
    assert _resolve_ref(0.12) is None
    assert _resolve_ref("普通文字") is None


def test_读取跨表引用的值(tmp_path):
    from openpyxl import load_workbook

    from ashare_verify import _cell_value
    p = _wb(tmp_path, "B5", "=1b_驱动分解!B5")
    wb = load_workbook(p, data_only=False)
    assert abs(_cell_value(wb, wb["1_假设"], "B1") - 0.1234) < 1e-12


def test_多层引用有上限(tmp_path):
    """防止循环引用把读取器转死。"""
    from openpyxl import load_workbook

    from ashare_verify import _cell_value
    wb = Workbook()
    a = wb.active
    a.title = "1_假设"
    a["B1"] = "=B2"
    a["B2"] = "=B1"          # 循环
    p = tmp_path / "loop.xlsx"
    wb.save(p)
    wb2 = load_workbook(p, data_only=False)
    with pytest.raises(ValueError, match="引用层数"):
        _cell_value(wb2, wb2["1_假设"], "B1")
