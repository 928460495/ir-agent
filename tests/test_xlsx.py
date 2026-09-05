"""Excel 导出。

核心立场: **导出活公式，不是死数字。**
若三张表和比率全是硬编码值，工作簿就是一张打印件 —— 改一个输入不会有任何
反应，勾稽也无法重算，「可复用」是空话。所以:
  · 比率、勾稽、可比全部用 Excel 公式引用三表单元格
  · 只有从数据源直接取得的 REPORTED 事实才写字面值
  · 每个字面值旁边保留 source_id 与披露日 —— 溯源不能在导出环节丢失
"""
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from ir_agent.ledger import Fact, FactLedger, Method
from ir_agent.xlsx import SHEETS, build_workbook

D = Decimal


def led() -> FactLedger:
    l = FactLedger()
    vals = dict(revenue=1000, cost_of_revenue=600, net_profit=200,
                net_profit_attr_parent=180, minority_interest_profit=20,
                total_assets=5000, total_liabilities=3000, total_equity=2000,
                equity_attr_parent=1800, minority_equity=200,
                cash_begin=300, cash_net_change=50, cash_end=350,
                cf_net_profit=200)
    for k, v in vals.items():
        l.put(Fact(key=k, value=D(str(v)), unit="元", currency="CNY",
                   period="2025FY", as_of=date(2026, 4, 20),
                   source_id="em_g_balance_x", method=Method.REPORTED))
    return l


@pytest.fixture
def wb(tmp_path):
    p = tmp_path / "t.xlsx"
    build_workbook(led(), code="600519", period="2025FY",
                   as_of=date(2026, 6, 1), path=p)
    return load_workbook(p)          # 公式形态


@pytest.fixture
def wbv(tmp_path):
    p = tmp_path / "t.xlsx"
    build_workbook(led(), code="600519", period="2025FY",
                   as_of=date(2026, 6, 1), path=p)
    return load_workbook(p, data_only=True)


def cells(ws):
    return {c.value for r in ws.iter_rows() for c in r if c.value is not None}


class TestStructure:
    def test_all_expected_sheets_exist(self, wb):
        assert set(SHEETS) <= set(wb.sheetnames)

    def test_metadata_sheet_records_the_as_of(self, wb):
        assert "2026-06-01" in {str(v) for v in cells(wb["说明"])}

    def test_metadata_records_the_analysed_period(self, wb):
        assert "2025FY" in {str(v) for v in cells(wb["说明"])}


class TestFormulasNotHardcodedValues:
    def test_gross_margin_is_a_formula(self, wb):
        f = _find(wb["比率"], "毛利率")
        assert isinstance(f, str) and f.startswith("=")

    def test_gross_margin_formula_references_the_income_sheet(self, wb):
        assert "利润表" in _find(wb["比率"], "毛利率")

    def test_reconciliation_check_is_a_formula(self, wb):
        f = _find(wb["勾稽校验"], "资产=负债+所有者权益")
        assert isinstance(f, str) and f.startswith("=")

    def test_reported_facts_are_written_as_literals(self, wb):
        """从数据源直接取得的数字必须是字面值，不能是公式 —— 它们是输入。"""
        v = _find(wb["利润表"], "营业收入")
        assert v == 1000

    def test_derived_rows_do_not_duplicate_source_numbers(self, wb):
        """净利率不能写成 0.2，必须由公式算出，否则改输入不会联动。"""
        assert str(_find(wb["比率"], "净利率")).startswith("=")


class TestProvenanceSurvivesExport:
    def test_source_id_is_carried_into_the_sheet(self, wb):
        assert any("em_g_balance_x" in str(v) for v in cells(wb["溯源"]))

    def test_disclosure_date_is_carried(self, wb):
        assert any("2026-04-20" in str(v) for v in cells(wb["溯源"]))

    def test_every_reported_fact_has_a_provenance_row(self, wb):
        n = sum(1 for r in wb["溯源"].iter_rows(min_row=2) if r[0].value)
        assert n >= 14


class TestReusability:
    def test_workbook_has_an_inputs_section_for_future_dcf(self, wb):
        """DCF 尚未接线，但假设区要先留好 —— 否则后来只能重做工作簿。"""
        assert "假设" in wb.sheetnames

    def test_assumptions_are_empty_but_labelled(self, wb):
        labels = {str(v) for v in cells(wb["假设"])}
        assert any("折现率" in s or "WACC" in s for s in labels)

    def test_no_sheet_is_completely_empty(self, wb):
        for name in SHEETS:
            assert cells(wb[name]), f"{name} 是空表"


def _find(ws, label):
    """返回标签所在行的第一个数据列的值。"""
    for row in ws.iter_rows():
        if row and str(row[0].value).strip() == label:
            for c in row[1:]:
                if c.value is not None:
                    return c.value
    return None


class TestFormulasActuallyEvaluate:
    """openpyxl 只写公式不求值 —— 引用写错也不会报错。
    这里解析公式引用的单元格、取值算一遍，与 Python 流水线独立算出的
    结果比对: 两条路径互证，比装一个 Excel 求值引擎更有力。"""

    def _eval(self, wb, sheet, label):
        import re
        f = _find(wb[sheet], label)
        assert isinstance(f, str) and f.startswith("="), f"{label} 不是公式"
        expr = f[1:]
        for m in sorted(set(re.findall(r"[一-龥]+![A-Z]+\d+", expr)),
                        key=len, reverse=True):
            sh, cell = m.split("!")
            v = wb[sh][cell].value
            assert v is not None, f"{m} 为空"
            expr = expr.replace(m, f"({v})")
        return eval(expr)                          # noqa: S307 受控表达式

    def test_gross_margin_matches_the_python_operator(self, wb):
        from ir_agent.operators import margins
        l = led()
        margins.compute(l, "2025FY", as_of=date(2026, 6, 1))
        py = float(l.get("gross_margin", "2025FY", as_of=date(2026, 6, 1)).value)
        assert self._eval(wb, "比率", "毛利率") == pytest.approx(py, rel=1e-9)

    def test_roe_matches_the_python_operator(self, wb):
        from ir_agent.operators import dupont
        l = led()
        dupont.compute(l, "2025FY", as_of=date(2026, 6, 1))
        py = float(l.get("roe", "2025FY", as_of=date(2026, 6, 1)).value)
        assert self._eval(wb, "比率", "ROE") == pytest.approx(py, rel=1e-6)

    def test_balance_check_evaluates_to_zero(self, wb):
        assert self._eval(wb, "勾稽校验", "资产=负债+所有者权益") == pytest.approx(0)

    def test_profit_split_check_evaluates_to_zero(self, wb):
        assert self._eval(wb, "勾稽校验", "净利润=归母+少数股东损益") == pytest.approx(0)

    def test_cash_check_evaluates_to_zero(self, wb):
        assert self._eval(wb, "勾稽校验", "期末现金=期初+净增加额") == pytest.approx(0)

    def test_a_broken_input_makes_the_check_non_zero(self, tmp_path):
        """把资产改掉，勾稽公式必须变成非零 —— 证明它真的在算，不是摆设。"""
        l = led()
        bad = Fact(key="total_assets", value=D("9999"), unit="元", currency="CNY",
                   period="2025FY", as_of=date(2026, 4, 21),
                   source_id="tampered", method=Method.REPORTED)
        l.put(bad)
        p = tmp_path / "b.xlsx"
        build_workbook(l, code="600519", period="2025FY",
                       as_of=date(2026, 6, 1), path=p)
        w = load_workbook(p)
        assert self._eval(w, "勾稽校验", "资产=负债+所有者权益") != pytest.approx(0)
