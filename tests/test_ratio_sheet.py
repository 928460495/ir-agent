# -*- coding: utf-8 -*-
"""7_财务比率与杜邦：中期列口径与负值显示。

**问题一**：中期列直接拿半年流量除以存量，得到的「净利率」还算对，但
「总资产周转率」「ROE」这类含流量÷存量的比率会系统性偏低约一半——读者
看到的是「今年效率腰斩」，实际只是少算了半年。

**问题二**：EBIT 利息保障倍数用 IFERROR 包着，EBIT 为负时返回空白，
读者看到的是「没有数据」，而真相是「利润不够付利息」——**最该看见的
那一格恰好被藏起来了**。
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pickle
from openpyxl import load_workbook
import ashare_build as B

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "cal_600183.pkl")


def _wb(tmp_path):
    cal = pickle.load(open(FIX, "rb"))
    out = str(tmp_path / "t.xlsx")
    B.build(cal, out)
    return load_workbook(out), cal


def test_中期列的流量类比率已年化(tmp_path):
    ws, _ = _wb(tmp_path)
    ws = ws["7_财务比率与杜邦"]
    hit = [c for row in ws.iter_rows() for c in row
           if c.value and isinstance(c.value, str) and "年化" in str(c.value)]
    assert hit, "中期列的流量类比率必须标明已年化，否则读者会当成全年口径"


def test_ROE中期列公式已按share倒数年化(tmp_path):
    """ROE 是「流量÷平均存量」，中期列须对分子（归母净利）乘以 1/share。

    fixture 的 share=0.4785 → 年化因子 1/0.4785≈2.0899，而不是固定的 2.0
    （固定 ×2 正是 `_wc` 注释里说会系统性带偏的做法）。
    """
    cal = pickle.load(open(FIX, "rb"))
    wb, _ = _wb(tmp_path)
    ws = wb["7_财务比率与杜邦"]
    row = [r for r in ws.iter_rows() if r[0].value == "ROE（归母，平均净资产）"]
    assert row, "找不到 ROE 行"
    R = cal["_rows"] if "_rows" in cal else None
    NH = 6  # 5 个年报 + 1 个中报，来自 fixture periods.labels
    itm_cell = row[0][1 + (NH - 1)]  # B=索引1 对应第0列，中期列是第 NH-1 列（0-based）
    f = itm_cell.value
    assert isinstance(f, str) and f.startswith("="), "ROE 中期列必须是活公式"
    factor = round(1.0 / cal["interim"]["share"], 6)
    assert str(factor) in f or f"{factor:.6f}" in f, (
        f"ROE 中期列公式未按 1/share={factor} 年化，公式为：{f}")
    # 年报列不应被年化——不能一刀切
    annual_cell = row[0][1]  # 第一个年报列
    af = annual_cell.value
    assert af and str(factor) not in af and f"{factor:.6f}" not in af, (
        "年报列的 ROE 不应被年化因子污染")


def test_DSO中期列分母已年化(tmp_path):
    """DSO 是「存量÷流量×365」，中期列须对分母（营收）乘以 1/share，
    而不是像旧实现那样直接拿半年营收算，导致天数虚高约一倍。
    """
    cal = pickle.load(open(FIX, "rb"))
    wb, _ = _wb(tmp_path)
    ws = wb["7_财务比率与杜邦"]
    row = [r for r in ws.iter_rows() if r[0].value == "应收周转天数 DSO"]
    assert row, "找不到 DSO 行"
    NH = 6
    itm_cell = row[0][1 + (NH - 1)]
    f = itm_cell.value
    factor = round(1.0 / cal["interim"]["share"], 6)
    assert isinstance(f, str) and (str(factor) in f or f"{factor:.6f}" in f), (
        f"DSO 中期列公式未按 1/share={factor} 年化分母，公式为：{f}")


def test_毛利率等流量比流量不年化(tmp_path):
    """毛利率、营业利润率、EBIT 利润率、归母净利率——本来就对，不能被一刀切年化。"""
    cal = pickle.load(open(FIX, "rb"))
    wb, _ = _wb(tmp_path)
    ws = wb["7_财务比率与杜邦"]
    factor = round(1.0 / cal["interim"]["share"], 6)
    for label in ("毛利率", "营业利润率", "EBIT 利润率", "归母净利率"):
        row = [r for r in ws.iter_rows() if r[0].value == label]
        assert row, f"找不到 {label} 行"
        NH = 6
        itm_cell = row[0][1 + (NH - 1)]
        f = itm_cell.value
        assert isinstance(f, str)
        assert str(factor) not in f and f"{factor:.6f}" not in f, (
            f"{label} 是流量÷流量，中期列不该被年化，公式为：{f}")


def test_利息保障倍数为负时不被藏起来(tmp_path):
    ws, _ = _wb(tmp_path)
    ws = ws["7_财务比率与杜邦"]
    row = [r for r in ws.iter_rows()
           if r[0].value == "EBIT 利息保障倍数"]
    assert row, "找不到利息保障倍数行"
    formulas = [c.value for c in row[0][1:] if c.value]
    assert formulas, "该行不能整行为空"
    assert not any('IFERROR' in f and '""' in f for f in formulas if isinstance(f, str)), \
        "IFERROR 回退成空字符串会把「利润不够付利息」这件事藏起来"
    assert any('利润不足' in f for f in formulas if isinstance(f, str)), \
        "应显式回退成「利润不足」，而不是让读者猜"
