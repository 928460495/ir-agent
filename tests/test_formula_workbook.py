"""跨仓库对账：工作簿「1b_驱动分解」的合成格 ≡ ir-agent 计算引擎（spec 2026-10-14 四）。

翻译只要错一个引用（上期取成本年、基期取成上期、档位列错位……），这里就对不上。
工作簿合成格用 `ashare_verify` 自己的公式引擎求值（与「公式引擎验算全过」同一套），
不用 ir-agent 的引擎去算工作簿 —— 两边各算各的，才叫对账。
"""
from __future__ import annotations

import os
import pickle
from decimal import Decimal as D
from pathlib import Path

import pytest

from ir_agent import assume, dcf
from ir_agent.assume import Case
from ir_agent.drivers import Decomposition, Factor, Op
from ir_agent.methodology import parse as parse_meth

_DIR = Path(os.environ.get(dcf.ENV_DIR) or dcf.DEFAULT_DIR)
_FIX = _DIR / "tests" / "fixtures" / "cal_600183.pkl"
try:
    _MODS = dcf._load(_DIR)
except dcf.DcfUnavailable:
    _MODS = None
pytestmark = pytest.mark.skipif(
    _MODS is None or not _FIX.exists(),
    reason="未找到 ashare-dcf 或其离线夹具 tests/fixtures/cal_600183.pkl")

Q = "「依据原文」"


def _cal():
    with open(_FIX, "rb") as f:
        return pickle.load(f)


def _rows(spec: dict, years: int) -> str:
    """spec: {短名: (悲观, 中性, 乐观)} → 逐年的假设行。取值可以是字符串（逐年
    同值），也可以是 `lambda y: str`（随年份变 —— 水平模式下水平逐年同值会让第
    2 年起三档营收增速都是 0，被合成后单调校验正确地驳回）。"""
    def v(x, y):
        return x(y) if callable(x) else x
    return "\n".join(
        f"假设: {k} | 第{y}年 | {v(lo, y)} | {v(mid, y)} | {v(hi, y)} | 理由 {Q}"
        for y in range(1, years + 1) for k, (lo, mid, hi) in spec.items())


LEVEL = parse_meth("""模式: 水平
因子: 量 | 覆铜板出货量 | 万平方米 | 年报产销量 | 产销量
因子: 价 | 覆铜板平均单价 | 元/平方米 | 研报价格跟踪 | 平均单价
因子: 本 | 覆铜板单位成本 | 元/平方米 | 年报成本构成 | 单位成本
公式: 营收 = 量 × 价
公式: 毛利率 = (价 − 本) ÷ 价
""").model
LEVEL_ROWS = {"量": tuple((lambda g: lambda y: f"{100 * (1 + g) ** y:.4f}")(g)
                          for g in (0.02, 0.05, 0.08)),
              "价": ("9", "10", "11"), "本": ("7.5", "7.4", "7.3")}
LEVEL_BASE = f"基期: 量 | 100 | {Q}\n基期: 价 | 9.5 | {Q}\n基期: 本 | 7.2 | {Q}"

CHANGE = parse_meth("""模式: 变化
因子: 稼动率变动 | 稼动率变动 | pp | 年报 | 产能利用率
因子: 价增 | 平均单价同比变化率 | 变化率 | 研报 | 平均单价
因子: 提价 | 提价对毛利率的逐年贡献 | pp | 调价函 | 调价
因子: 原料 | 原材料对毛利率的逐年侵蚀 | pp | 年报 | 原材料价格
常量: 上期稼动率 | 78% | 年报经营情况讨论「产能利用率约78%」
公式: 营收 = 上期(营收) × (1 + 稼动率变动 / 100 / 上期稼动率) × (1 + 价增)
公式: 毛利率 = 上期(毛利率) + (提价 + 原料) / 100
""").model
CHANGE_ROWS = {"稼动率变动": ("-2", "1", "3"), "价增": ("-3%", "1%", "4%"),
               "提价": ("-0.5", "0.3", "0.8"), "原料": ("-0.6", "-0.2", "0.1")}

LEGACY = [Decomposition("revenue", Op.MULTIPLY, [Factor("cap", "产能", "变化率"),
                                                  Factor("asp", "单价", "变化率")]),
          Decomposition("gross_margin", Op.ADD_PP, [Factor("mix", "结构", "pp"),
                                                     Factor("cost", "成本", "pp")])]
LEGACY_ROWS = {"产能": ("2%", "5%", "8%"), "单价": ("-3%", "0%", "2%"),
               "结构": ("-1", "0.5", "1.5"), "成本": ("-2", "-1", "-0.5")}


def _build(model, rows, extra, tmp_path):
    D_, C_, B_, V_ = _MODS
    cal = _cal()
    years = len(cal["periods"]["proj"])
    a = assume.parse(_rows(rows, years) + ("\n" + extra if extra else ""), years,
                     decomps=model)
    a.labels = list(cal["periods"]["proj"])
    out = assume.apply_to(cal, a)
    out["judg"]["drivers"] = assume.to_driver_block(a, cal)
    path = tmp_path / "w.xlsx"
    B_.build(out, str(path))
    return a, cal, path


def _sheet_values(path, label_suffix, cal):
    from openpyxl import load_workbook

    V_ = _MODS[3]
    wb = load_workbook(path)
    ds = wb["1b_驱动分解"]
    rowmap = {str(ds.cell(r, 1).value): r for r in range(1, ds.max_row + 1)}
    out = {}
    for col, case in zip("BCD", ("bull", "base", "bear")):
        out[case] = [D(str(V_._cell_value(wb, ds, f"{col}{rowmap[f'{y} {label_suffix}']}")))
                     for y in cal["periods"]["proj"]]
    return out, wb, ds, rowmap


def _engine(a, cal):
    bm = assume._base_margin_of(cal)
    bases = assume.margin_bases(a, bm, assume.anchor_from_cal(cal))
    return {c.name.lower(): assume.synthesize(a, c, bm, bases)
            for c in (Case.BULL, Case.BASE, Case.BEAR)}


@pytest.mark.parametrize("model,rows,extra", [
    (LEVEL, LEVEL_ROWS, LEVEL_BASE),
    (CHANGE, CHANGE_ROWS, ""),
    (LEGACY, LEGACY_ROWS, ""),
], ids=["水平模式", "变化模式_常量与上期毛利率", "旧乘加分解式"])
def test_工作簿合成格与ir_agent引擎逐档逐年一致(model, rows, extra, tmp_path):
    a, cal, path = _build(model, rows, extra, tmp_path)
    want = _engine(a, cal)
    growth, *_ = _sheet_values(path, "营收增速（合成）", cal)
    margin, *_ = _sheet_values(path, "毛利率（合成）", cal)
    for case in ("bull", "base", "bear"):
        for y in range(len(growth[case])):
            assert abs(growth[case][y] - want[case][0][y]) < D("1e-9"), (case, y)
            assert abs(margin[case][y] - want[case][1][y]) < D("1e-9"), (case, y)


def test_工作簿公式引擎验算全过(tmp_path):
    _, _, path = _build(LEVEL, LEVEL_ROWS, LEVEL_BASE, tmp_path)
    fails, _ = _MODS[3].verify(str(path), quiet=True)
    assert not fails


def test_常量区与基期区存在_改一个因子合成格随之变化(tmp_path):
    _, cal, path = _build(CHANGE, CHANGE_ROWS, "", tmp_path)
    before, wb, ds, rowmap = _sheet_values(path, "营收增速（合成）", cal)
    assert "常量 上期稼动率" in rowmap and "基期 营收（亿元）" in rowmap
    y1 = cal["periods"]["proj"][0]
    ds[f"C{rowmap[f'{y1} 平均单价同比变化率']}"] = 0.20
    V_ = _MODS[3]
    after = D(str(V_._cell_value(wb, ds, f"C{rowmap[f'{y1} 营收增速（合成）']}")))
    assert after != before["base"][0]
