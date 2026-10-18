# -*- coding: utf-8 -*-
"""工作簿「1_假设」第三小节：首个预测年按季度估算。

锚定开关打开后，首个预测年的营收与毛利率取季度合计：已公布季度用实际，
未公布季度 = 去年同季 × (1 + 年初至今同比)，营收与营业成本各用各的同比。
首年三档相同。规则本身见 ashare_calibrate.anchor_quarterly 与
tests/test_anchor_quarterly.py；这里测工作簿的活公式、独立重算与情景切换。

（2026-10-18 之前这一节是「已披露 ÷ 近几年同期占比」的三档锚定，对周期股
失灵——牧原股份首年营收高估约 13.6%。）
"""
import copy, os, pickle, sys
import pytest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook
import ashare_build as B
import ashare_calibrate as C
import ashare_verify as V

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "cal_600183.pkl")
SCEN_LABEL = "选择情景  1=Bull乐观 / 2=Base中性 / 3=Bear悲观"
PFX = "季度估算："


@pytest.fixture(scope="module")
def cal():
    """夹具快照的中期块由当前 calibrate（C.interim_block）从夹具自带的报表生成——
    600183 的真实季度数据（2026 上半年营收同比 +50.0%）。"""
    if not os.path.exists(FIX):
        pytest.skip("缺离线夹具：python3 tests/make_fixture.py 600183")
    with open(FIX, "rb") as f:
        c = pickle.load(f)
    assert c["interim"] and c["interim"].get("fy_rev"), "夹具中期块须是新格式"
    return c


@pytest.fixture(scope="module")
def cal_old(cal):
    """旧版 calibrate 的中期块形状（缓存的 pickle）：只有累计实际与历史占比。"""
    c = copy.deepcopy(cal)
    it = c["interim"]
    c["interim"] = dict(period=it["period"], label=it["label"], rev=it["rev"],
                        cogs=it["cogs"], gm=it["gm"], share=0.4785,
                        share_hist=[0.5204, 0.4751, 0.4723, 0.446], share_std=0.0267)
    return c


def _rowmap(ws):
    m = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip():
            m.setdefault(v.strip(), r)
    return m


def _build(cal, tmp_path, name="w.xlsx"):
    p = tmp_path / name
    B.build(copy.deepcopy(cal), str(p))
    return p


def test_已公布季度是实际值_未公布季度是外推公式(cal, tmp_path):
    a = load_workbook(_build(cal, tmp_path))["1_假设"]
    m = _rowmap(a)
    it = cal["interim"]
    cur, prev = m[PFX + "本年营业总收入（单季）"], m[PFX + "去年营业总收入（单季）"]
    g = m[PFX + "营收年初至今同比"]
    for j in range(2):                       # Q1、Q2：实际值
        assert a.cell(cur, 2 + j).value == pytest.approx(it["q_cur_rev"][j], abs=1e-4)
    for j in (2, 3):                         # Q3、Q4：去年同季 × (1 + 同比)
        f = str(a.cell(cur, 2 + j).value)
        col = "BCDE"[j]
        assert f.startswith("=") and f"{col}{prev}" in f and f"$B${g}" in f, f
    for j in range(4):
        assert a.cell(prev, 2 + j).value == pytest.approx(it["q_prev_rev"][j], abs=1e-4)


def test_年初至今同比是公式_营收与成本各一行(cal, tmp_path):
    a = load_workbook(_build(cal, tmp_path))["1_假设"]
    m = _rowmap(a)
    for item in ("营收", "营业成本"):
        f = str(a.cell(m[f"{PFX}{item}年初至今同比"], 2).value)
        assert f.startswith("=SUM(") and "-1" in f.replace(" ", ""), f


def test_隐含全年等于本年四季合计(cal, tmp_path):
    a = load_workbook(_build(cal, tmp_path))["1_假设"]
    m = _rowmap(a)
    cur_r, cur_c = m[PFX + "本年营业总收入（单季）"], m[PFX + "本年营业成本（单季）"]
    assert str(a.cell(m["→ 隐含全年营业总收入"], 2).value) == f"=SUM(B{cur_r}:E{cur_r})"
    assert str(a.cell(m["→ 隐含全年营业成本"], 2).value) == f"=SUM(B{cur_c}:E{cur_c})"


def test_独立重算的首年与calibrate一致(cal, tmp_path):
    it = V.read_inputs(str(_build(cal, tmp_path)))["it"]
    assert it["fy_rev"] == pytest.approx(cal["interim"]["fy_rev"], rel=1e-6)
    assert it["fy_cogs"] == pytest.approx(cal["interim"]["fy_cogs"], rel=1e-6)


def test_首年三档相同_切换情景首年生效值不变(cal, tmp_path):
    p = _build(cal, tmp_path)
    firsts = []
    for scen in (1, 2, 3):
        wb = load_workbook(p)
        a = wb["1_假设"]
        a.cell(_rowmap(a)[SCEN_LABEL], 2).value = scen
        wb.save(p)
        rows, _ = V.recompute(V.read_inputs(str(p)))
        firsts.append((round(rows[0]["rev"], 6), round(rows[0]["cogs"], 6)))
    assert firsts[0] == firsts[1] == firsts[2], firsts
    assert firsts[0][0] == pytest.approx(cal["interim"]["fy_rev"], rel=1e-6)
    assert firsts[0][1] == pytest.approx(cal["interim"]["fy_cogs"], rel=1e-6)


def test_公式引擎验算仍然全过(cal, tmp_path):
    p = _build(cal, tmp_path)
    fails, _ = V.verify(str(p), quiet=True)
    assert not fails
    bad = [x for x in V.verify_all_scenarios(str(p), quiet=True) if x[0] != "Base"]
    assert not bad


def test_旧格式中期块不做首年估算(cal_old, tmp_path):
    """旧版 calibrate 的产出（缓存的 pickle）没有季度明细，无从估算：不建第三小节，
    首年回到增速判断项——不拿旧的三档规则凑一个数。"""
    a = load_workbook(_build(cal_old, tmp_path, "old.xlsx"))["1_假设"]
    m = _rowmap(a)
    assert "→ 隐含全年营业总收入" not in m
    assert not any(k.startswith(PFX) for k in m)


# ── I2（final-fix-D 终审）：Bull/Bear 用的 σ 不该混进中报 ────────────────
#
# 首年三档锚定取消后 σ 不再用于锚定，但 gm_stats 仍用它；口径测试保留。

_GM_HIST_600183 = [0.2682, 0.2203, 0.1919, 0.2204, 0.2647, 0.3069]
_ANN_600183 = ["20211231", "20221231", "20231231", "20241231", "20251231"]
_REV_600183 = {"20211231": 202.74262977869998, "20221231": 180.14442222419999,
              "20231231": 165.860728088, "20241231": 203.88330213560002,
              "20251231": 284.3113845944, "20260630": 190.2561176613}
_COGS_600183 = {p: 1.0 for p in _ANN_600183 + ["20260630"]}  # 只需非空即可，n_ann_gm 只判存在性


def test_gm_std按年度序列算不混进中报():
    """**600183 实跑复现。** `np.std(gm_hist)`（此前实际用于三档的口径）与
    `np.std(gm_hist[:5])`（F 列文案声称的"历史年度毛利率标准差"）差 31%，
    只因为 gm_hist 末项是中报，混进了不该混的一项。"""
    import numpy as np
    import ashare_calibrate as C
    sigma_full = round(float(np.std(_GM_HIST_600183)), 4)
    sigma_ann = round(float(np.std(_GM_HIST_600183[:5])), 4)
    assert sigma_full == 0.0383, "混进中报的（此前的）口径"
    assert sigma_ann == 0.0292, "只按年度序列算的（文案声称的）口径"
    got = C._gm_std_annual(_GM_HIST_600183, _ANN_600183, _REV_600183, _COGS_600183)
    assert round(got, 4) == 0.0292, "calibrate 实际用的 σ 必须是年度口径，不是混了中报的那个"


def test_gm_std_annual与gm_mean用同一套过滤判据():
    """`gm_mean` 已经用 `n_ann_gm`（数 rev 与 cogs 都有的年度期数）过滤掉
    中报——近 5 年缺 1 年营业成本时，年度项数会小于 len(ann)，此时 σ 也
    必须跟着少算这一年，不能自己另开一个"取前 len(ann) 项"的口径。"""
    import ashare_calibrate as C
    ann = _ANN_600183
    rev = dict(_REV_600183)
    cogs = {p: 1.0 for p in ann[1:] + ["20260630"]}   # 缺最早一年的营业成本
    gm_hist = [0.20, 0.21, 0.22, 0.23, 0.24, 0.30]     # 只有 4 个年度项 + 1 个中报
    got = C._gm_std_annual(gm_hist, ann, rev, cogs)
    import numpy as np
    assert got == pytest.approx(float(np.std(gm_hist[:4])), abs=1e-9)


# 旧乘/加分解式经 ir-agent `formula.from_decomps` 翻译后的两条公式的记号序列
# （`Formula.tokens()`，全括号）：营收 = 上期(营收)*(1+产能)*(1+单价)；
# 毛利率 = 基期(毛利率)+(结构+成本)/100。
_REV_TOKENS = [["lp", "("], ["lp", "("], ["prev", "营收"], ["op", "*"],
               ["lp", "("], ["num", "1"], ["op", "+"], ["name", "产能"], ["rp", ")"],
               ["rp", ")"], ["op", "*"],
               ["lp", "("], ["num", "1"], ["op", "+"], ["name", "单价"], ["rp", ")"],
               ["rp", ")"]]
_GM_TOKENS = [["lp", "("], ["base", "毛利率"], ["op", "+"],
              ["lp", "("], ["lp", "("], ["name", "结构"], ["op", "+"],
              ["name", "成本"], ["rp", ")"], ["op", "/"], ["num", "100"], ["rp", ")"],
              ["rp", ")"]]


def _with_drivers(cal, base_levels=None, anchor_y1=None):
    """给夹具 cal 注入一份最小的因子分解块（ir-agent 的 to_driver_block 产出形状，
    schema 2：一个块、公式以记号交付）。"""
    c = copy.deepcopy(cal)
    n = len(c["periods"]["proj"])
    annual = c["periods"]["annual"]
    drv = {"schema": 2, "mode": "变化",
           "factors": [{"key": "cap", "short": "产能", "label": "产能", "unit": "变化率"},
                       {"key": "asp", "short": "单价", "label": "单价", "unit": "变化率"},
                       {"key": "mix", "short": "结构", "label": "结构", "unit": "pp"},
                       {"key": "cost", "short": "成本", "label": "成本", "unit": "pp"}],
           "constants": [],
           "formulas": {"revenue": _REV_TOKENS, "gross_margin": _GM_TOKENS},
           "formula_text": {"revenue": "营收 = 上期(营收)*(1+产能)*(1+单价)",
                            "gross_margin": "毛利率 = 基期(毛利率)+(结构+成本)/100"},
           "base": {"revenue": c["rev"][annual[-1]], "margin": 0.2647, "factors": {}},
           "values": {"bull": [[0.05, 0.01, 1.0, 0.5]] * n,
                      "base": [[0.05, 0.01, 0.5, 0.2]] * n,
                      "bear": [[0.05, 0.01, 0.0, -0.2]] * n}}
    if base_levels:
        drv["base_levels"] = base_levels
    if anchor_y1:
        drv["anchor_y1"] = anchor_y1
    c["judg"]["drivers"] = drv
    return c


def _gm_synth_formula(path, year_label, col):
    ds = load_workbook(path)["1b_驱动分解"]
    r = _rowmap(ds)[f"{year_label} 毛利率（合成）"]
    return str(ds[f"{col}{r}"].value)


def test_1b毛利率合成可以按档按年用不同基期(cal, tmp_path):
    """锚定开启时，第 2 年起的毛利率要从各档锚定后的首年水平接着走，而不是回到
    2025 年的基期——生益科技实跑 Base 档 2026 年锚定 30.69%，2027 年却掉回
    25.87%。基期因此要能按档、按年给。"""
    n = len(cal["periods"]["proj"])
    bl = {"bull": [0.2647] + [0.3010] * (n - 1),
          "base": [0.2647] + [0.2900] * (n - 1),
          "bear": [0.2647] + [0.2800] * (n - 1)}
    p = tmp_path / "w.xlsx"
    B.build(_with_drivers(cal, bl), str(p))
    y1, y2 = cal["periods"]["proj"][0], cal["periods"]["proj"][1]
    assert _gm_synth_formula(p, y1, "B").startswith("=(0.2647+")
    assert _gm_synth_formula(p, y2, "B").startswith("=(0.301+")
    assert _gm_synth_formula(p, y2, "C").startswith("=(0.29+")
    assert _gm_synth_formula(p, y2, "D").startswith("=(0.28+")


def test_不给按年基期时维持单一基期(cal, tmp_path):
    p = tmp_path / "w.xlsx"
    B.build(_with_drivers(cal), str(p))
    y2 = cal["periods"]["proj"][1]
    ds = load_workbook(p)["1b_驱动分解"]
    gm0 = _rowmap(ds)["基期 毛利率"]
    for col in "BCD":
        assert _gm_synth_formula(p, y2, col).startswith(f"=($B${gm0}+")


def test_按年基期之下公式引擎验算仍然全过(cal, tmp_path):
    n = len(cal["periods"]["proj"])
    bl = {k: [0.2647] + [v] * (n - 1)
          for k, v in (("bull", 0.301), ("base", 0.29), ("bear", 0.28))}
    p = tmp_path / "w.xlsx"
    B.build(_with_drivers(cal, bl), str(p))
    fails, _ = V.verify(str(p), quiet=True)
    assert not fails


def _level_block(cal):
    """水平模式：营收 = 量 × 价；毛利率 = (价 − 本) ÷ 价；带一个常量。"""
    c = copy.deepcopy(cal)
    n = len(c["periods"]["proj"])
    lp, rp = ["lp", "("], ["rp", ")"]
    c["judg"]["drivers"] = {
        "schema": 2, "mode": "水平",
        "factors": [{"key": "q", "short": "量", "label": "出货量", "unit": "万吨"},
                    {"key": "p", "short": "价", "label": "单价", "unit": "元"},
                    {"key": "c", "short": "本", "label": "单位成本", "unit": "元"}],
        "constants": [{"name": "系数", "value": 1.0, "source": "示意"}],
        "formulas": {
            "revenue": [lp, lp, ["name", "量"], ["op", "*"], ["name", "价"], rp,
                        ["op", "*"], ["name", "系数"], rp],
            "gross_margin": [lp, lp, ["name", "价"], ["op", "-"], ["name", "本"], rp,
                             ["op", "/"], ["name", "价"], rp]},
        "base": {"revenue": None, "margin": 0.25,
                 "factors": {"量": 100.0, "价": 9.5, "本": 7.2}},
        "values": {"bull": [[100.0 * 1.08 ** (y + 1), 11.0, 7.3] for y in range(n)],
                   "base": [[100.0 * 1.05 ** (y + 1), 10.0, 7.4] for y in range(n)],
                   "bear": [[100.0 * 1.02 ** (y + 1), 9.0, 7.5] for y in range(n)]}}
    return c


def test_水平模式的基期营收由公式算出_常量区可引用(cal, tmp_path):
    p = tmp_path / "w.xlsx"
    B.build(_level_block(cal), str(p))
    ds = load_workbook(p)["1b_驱动分解"]
    m = _rowmap(ds)
    assert str(ds.cell(m["基期 营收（亿元）"], 2).value).startswith("=((")
    assert "常量 系数" in m and ds.cell(m["常量 系数"], 6).value == "示意"
    y1 = cal["periods"]["proj"][0]
    assert f"$B${m['常量 系数']}" in str(ds[f"C{m[f'{y1} 营收（合成）']}"].value)


def test_水平模式公式引擎验算全过(cal, tmp_path):
    p = tmp_path / "w.xlsx"
    B.build(_level_block(cal), str(p))
    fails, _ = V.verify(str(p), quiet=True)
    assert not fails


def test_未锚定时第三小节与比率表都写明原因(cal, tmp_path):
    """终审（2026-10-19）：首年不锚定此前是静默的，比率表还会写错原因
    （「缺去年同期数据」）。calibrate 把原因记进 interim_skip，工作簿照写。"""
    c = copy.deepcopy(cal)
    c["interim"] = None
    c["interim_skip"] = "去年 2025 Q3 营业成本缺失"
    wb = load_workbook(_build(c, tmp_path, "skip.xlsx"))
    a = wb["1_假设"]
    sec = [v for v in (a.cell(r, 1).value for r in range(1, a.max_row + 1))
           if isinstance(v, str) and v.startswith("三、首个预测年未锚定")]
    assert sec and "去年 2025 Q3 营业成本缺失" in sec[0], sec
    ws = wb["7_财务比率与杜邦"]
    texts = " ".join(str(ws.cell(r, c_).value) for r in range(1, 6) for c_ in range(1, 4))
    assert "去年 2025 Q3 营业成本缺失" in texts and "缺去年同期数据" not in texts


def test_季度表数据损坏时校验明确报错(cal, tmp_path):
    """季度表在、却算不出首年：静默放过会让校验悄悄不锚定，与工作簿对不上。"""
    p = _build(cal, tmp_path, "bad.xlsx")
    wb = load_workbook(p)
    a = wb["1_假设"]
    a.cell(_rowmap(a)[PFX + "去年营业成本（单季）"], 3).value = None
    wb.save(p)
    with pytest.raises(ValueError, match="季度估算"):
        V.read_inputs(str(p))


def test_季度表数值齐全但算不出首年时校验明确报错(cal, tmp_path):
    """格子都是数，但按规则不能估算（去年某季营业成本被改成负数）：同样不能静默
    不锚定——工作簿开关 = 1 时照样锚定，校验必须对得上或明确报错。"""
    p = _build(cal, tmp_path, "neg.xlsx")
    wb = load_workbook(p)
    a = wb["1_假设"]
    a.cell(_rowmap(a)[PFX + "去年营业成本（单季）"], 3).value = -1.0
    wb.save(p)
    with pytest.raises(ValueError, match="算不出首年.*不为正"):
        V.read_inputs(str(p))
