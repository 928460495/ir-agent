# -*- coding: utf-8 -*-
"""变异测试：故意破坏融资模块的每一条接线，确认验证器会报错（否则断言是摆设）。"""
import os, sys, shutil
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook
import ashare_verify as V
SRC=os.environ.get("DCF_TEST_XLSX") or sys.argv[1]   # 传入一个已生成的工作簿
def rowmap(ws):
    m={}
    for r in range(1,ws.max_row+1):
        v=ws.cell(r,1).value
        if isinstance(v,str) and v.strip(): m.setdefault(v.strip(),r)
    return m
ED={"增发股数（亿股）":1.1818,"发行价（元/股）":55.0,"到账年份序号（1=首个预测年）":1,
    "投向 — 资本开支（亿元）":40.0,"投向 — 偿还债务（亿元）":15.0,"募投资本开支建设期（年）":3}
os.makedirs("/tmp/dcf_fin_test",exist_ok=True)

def mut(tag, fn):
    dst=f"/tmp/dcf_fin_test/_m_{tag}.xlsx"; shutil.copy(SRC,dst)
    wb=load_workbook(dst); a=wb["1_假设"]; m=rowmap(a)
    for k,v in ED.items(): a.cell(m[k],2).value=v
    fn(wb); wb.save(dst)
    fails,_=V.verify(dst,quiet=True)
    ok = (len(fails)==0) if tag.startswith("00_") else (len(fails)>0)
    print(f"{tag:34s} → {len(fails):3d} 项失败  {'✅' if ok else '❌'}")
    if fails: print("      例:",fails[0])

mut("00_不破坏（应 0 失败）", lambda wb: None)
mut("01_权益漏加募资", lambda wb: wb["3_资产负债表"].__setitem__("H25","='3_资产负债表'!G25+'2_利润表'!H23+'4_现金流量表'!H16"))
mut("02_资本开支漏募投", lambda wb: wb["4_现金流量表"].__setitem__("H14","=-'2_利润表'!H6*'1_假设'!$E$30"))
mut("03_筹资净额漏募资", lambda wb: wb["4_现金流量表"].__setitem__("H19","=H16+H17+H18+H27+H28"))
mut("04_短贷漏定向偿还", lambda wb: wb["3_资产负债表"].__setitem__("H18","='3_资产负债表'!G18+'4_现金流量表'!H18"))
mut("05_DCF股本漏增发", lambda wb: (lambda w,m: w.cell(m["总股本（亿股，含增发）"],2,"='1_假设'!B99"))(wb["5_DCF"],rowmap(wb["5_DCF"])))
mut("06_权益桥漏募资净额", lambda wb: (lambda w,m: w.cell(m["归属全体股东权益价值"],2,
        f"=B{m['企业价值 EV（滚动至评估日）']}-B{m['减：净负债（基期）']}+B{m['加：非经营资产']}"))(wb["5_DCF"],rowmap(wb["5_DCF"])))
mut("07_EPS漏摊薄", lambda wb: wb["2_利润表"].__setitem__("H24","=H23/'1_假设'!$B$53"))
