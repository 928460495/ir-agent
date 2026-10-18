# -*- coding: utf-8 -*-
"""A股 DCF 模板 — 蒙特卡洛模拟层。

设计取舍（有实测依据，见 README）：
  · 计算在 Python（10 万路径 18ms；Excel 原生需 56 万个公式格且会打断验算体系）
  · 参数存在 Excel「8_蒙特卡洛」页，用户可编辑
  · 固定随机种子 → 结果完全可复现，可被断言

往返用法：
    python3 ashare_dcf.py 600183 --mc        # 建模型 + 蒙特卡洛页（默认参数）
    python3 ashare_mc.py <xlsx>              # 读该文件参数区，重跑，写回同一文件
"""
from __future__ import annotations
import sys, os, math, warnings
warnings.filterwarnings("ignore")
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

SHEET = "8_蒙特卡洛"
NAVY="1F4E79"; LBLUE="D9E1F2"; GREY="F2F2F2"; YEL="FFF2CC"; ORG="FCE4D6"
F_HDR=Font(color="FFFFFF",bold=True,size=11); F_COL=Font(bold=True,size=10)
F_IN=Font(color="0000FF",size=10); F_FX=Font(color="000000",size=10)
P_N=PatternFill("solid",fgColor=NAVY); P_B=PatternFill("solid",fgColor=LBLUE)
P_G=PatternFill("solid",fgColor=GREY); P_Y=PatternFill("solid",fgColor=YEL); P_O=PatternFill("solid",fgColor=ORG)
BX=Border(bottom=Side(style="thin",color="BFBFBF"))
NUM="#,##0.00"; PCT="0.0%"; PCT2="0.00%"; CNY="¥#,##0.00"; INT="#,##0"

# 参数表：(键, 标签, 默认值, 格式, 说明)
PARAM_SPEC = [
    ("n_paths",   "模拟路径数",                 50000, INT,  "1,000–1,000,000。10 万路径约 20ms"),
    ("seed",      "随机种子",                      42, "0",  "固定种子 → 结果可复现；改它得到另一组抽样"),
    ("sd_g",      "营收增速 路径平移 标准差",     None, PCT2, "持续性成分：默认 = 历史年度σ × 0.5（占总方差 25%）"),
    ("sd_g_yr",   "营收增速 逐年扰动 标准差",     None, PCT2, "暂时性成分：默认 = 历史年度σ × 0.866（占总方差 75%）"),
    ("lo_g",      "营收增速 平移 下限",          -0.30, PCT2, "截断边界，默认 ±2.5σ"),
    ("hi_g",      "营收增速 平移 上限",           0.30, PCT2, ""),
    ("sd_m",      "毛利率 路径平移 标准差",       None, PCT2, "持续性成分：默认 = 历史年度σ × 0.5"),
    ("sd_m_yr",   "毛利率 逐年扰动 标准差",       None, PCT2, "暂时性成分：默认 = 历史年度σ × 0.866"),
    ("lo_m",      "毛利率 平移 下限",            -0.12, PCT2, ""),
    ("hi_m",      "毛利率 平移 上限",             0.12, PCT2, ""),
    ("rho",       "ρ（增速 Δ 与 毛利率 Δ 的相关系数）", None, "0.000",
                  "涨价驱动型为正：增速高的年份毛利率也高。默认取历史值"),
    ("sd_beta",   "Beta 标准差",                  0.25, "0.000", "Beta → Ke → WACC；比直接抽 WACC 更有经济含义"),
    ("lo_beta",   "Beta 下限",                    0.40, "0.000", ""),
    ("hi_beta",   "Beta 上限",                    3.00, "0.000", ""),
    ("sd_gt",     "永续增长率 g 标准差",         0.008, PCT2, ""),
    ("lo_gt",     "永续增长率 g 下限",            0.00, PCT2, ""),
    ("gap_gt",    "g 与 WACC 的最小间距",        0.015, PCT2, "逐路径强制 g ≤ WACC − 此值，避免终值公式退化"),
    ("sd_cx",     "资本开支率 平移 标准差",      0.015, PCT2, ""),
    ("lo_cx",     "资本开支率 平移 下限",        -0.05, PCT2, ""),
    ("hi_cx",     "资本开支率 平移 上限",         0.05, PCT2, ""),
    ("sd_tax",    "实际税率 标准差",              0.00, PCT2,
                  "默认 0 = 税率固定为报告值、不参与随机化。若要开启，填入标准差即可"),
    ("lo_tax",    "实际税率 下限",                0.05, PCT2, ""),
    ("hi_tax",    "实际税率 上限",                0.25, PCT2, ""),
]

def _rowmap(ws):
    m={}
    for r in range(1,ws.max_row+1):
        v=ws.cell(r,1).value
        if isinstance(v,str) and v.strip(): m.setdefault(v.strip(), r)
    return m

# 方差分解：历史年度标准差 σ 拆成【持续性平移】与【逐年扰动】两部分，
# 满足 σ_shift² + σ_year² = σ²。持续性占 25% 方差 → σ_shift=0.5σ，σ_year=√0.75σ≈0.866σ。
# 直接把年度 σ 当作 5 年持续平移会严重高估离散度（实测标准差会超过中位数）。
PERSIST_SHARE = 0.25

def defaults(hist: dict) -> dict:
    d={k:v for k,_,v,_,_ in PARAM_SPEC}
    ks=math.sqrt(PERSIST_SHARE); ky=math.sqrt(1-PERSIST_SHARE)
    sg=hist.get("sd_growth",0.15); sm=hist.get("sd_margin",0.03)
    d["sd_g"]=round(sg*ks,4);  d["sd_g_yr"]=round(sg*ky,4)
    d["sd_m"]=round(sm*ks,4);  d["sd_m_yr"]=round(sm*ky,4)
    d["lo_g"]=round(-2.5*d["sd_g"],4); d["hi_g"]=round(2.5*d["sd_g"],4)
    d["lo_m"]=round(-2.5*d["sd_m"],4); d["hi_m"]=round(2.5*d["sd_m"],4)
    # 税率/Beta 边界也做成 σ 相对，避免固定边界离均值太近（生益税率 10.4%、σ=3%，
    # 固定下限 5% 只有 1.8σ → 实测 5.5% 拒绝率，且截断不对称会抬高有效均值）
    tx=hist.get("tax_mean"); bt=hist.get("beta_mean")
    if tx:
        # 税率默认不随机化（sd=0）。以下边界仅在用户手动开启随机化时才起作用。
        d["lo_tax"]=round(max(0.0, tx-0.06),4); d["hi_tax"]=round(min(0.25, tx+0.06),4)
    if bt: d["lo_beta"]=round(max(0.2, bt-2.5*d["sd_beta"]),3); d["hi_beta"]=round(bt+2.5*d["sd_beta"],3)
    d["rho"]=round(hist.get("rho",0.0),3)
    return d


def simulate(inp: dict, P: dict) -> dict:
    """向量化蒙特卡洛。与确定性模型逐项对应；所有标准差置 0 时必须复现基准。"""
    o=inp["o"]; G=inp["gen"]; D=inp["dcf"]; ny=inp["ny"]; IT=inp.get("it")
    F=inp.get("fin") or dict(ns=0.0,net=0.0,ucx=0.0,by=1,yr=1,disc=0.0)
    CXI=G.get("资本开支投向无形/长期待摊的比重",0.0)
    IPS=G.get("非现金减值减记中固定资产占比",1.0)
    n=int(P["n_paths"]); rng=np.random.default_rng(int(P["seed"]))

    g_live=np.array(inp["growth"],dtype=float); m_live=np.array(inp["gm"],dtype=float)
    cx_live=np.array(inp["capex"],dtype=float)
    anchored = bool(IT and IT["sw"]==1)
    if anchored:                       # 首年被已披露中期锁定（与 verify.recompute 一致）
        g_live[0]=IT["fy_rev"]/o["基期营业总收入"]-1
        m_live[0]=1-IT["fy_cogs"]/IT["fy_rev"]

    # --- 相关抽样：Cholesky 注入 ρ（增速平移 与 毛利率平移）---
    rho=float(np.clip(P["rho"],-0.999,0.999))
    L=np.linalg.cholesky(np.array([[1.0,rho],[rho,1.0]]))
    # 截尾用【拒绝重抽】而非 np.clip —— clip 会把越界样本堆在边界两点上形成点质量，
    # 恰好扭曲最需要准确的尾部（实测 1.25% 的增速Δ 被堆在 ±2.5σ 两点）。
    dg, dm, rej_gm = _rejection_pair(rng, L, P["sd_g"], P["sd_m"],
                                     P["lo_g"], P["hi_g"], P["lo_m"], P["hi_m"], n)
    eg, em = _corr_noise(rng, L, P["sd_g_yr"], P["sd_m_yr"], ny, n)   # 逐年扰动同样带相关性
    beta,rej_b = _rejection(rng, D["Beta（5年"], P["sd_beta"], P["lo_beta"], P["hi_beta"], n)
    tax,rej_t  = _rejection(rng, G["实际所得税率"], P["sd_tax"], P["lo_tax"], P["hi_tax"], n)
    dcx,rej_c  = _rejection(rng, 0.0, P["sd_cx"], P["lo_cx"], P["hi_cx"], n)
    ke=D["无风险利率 Rf"]+beta*D["股权风险溢价 ERP"]
    W=(1-D["债务权重 D/(D+E)"])*ke + D["债务权重 D/(D+E)"]*D["债务成本 Kd"]*(1-tax)
    gt=np.clip(rng.normal(D["永续增长率 g"],P["sd_gt"],n), P["lo_gt"], None)
    gt_raw=gt.copy()
    gt=np.minimum(gt, W-P["gap_gt"])                      # 逐路径保证 WACC−g 有下界
    rej_g=float((gt_raw>gt).mean())
    clipped=float(max(rej_gm, rej_b, rej_t, rej_c, rej_g))
    rej_detail=dict(增速与毛利率=rej_gm, Beta=rej_b, 实际税率=rej_t, 资本开支=rej_c, 永续g=rej_g)

    opex=G["税金及附加率"]+G["销售费用率"]+G["管理费用率"]+G["研发费用率"]
    oth=G["其他收益 / 营收"]+G["其他经营损益 / 营收"]
    rev=np.full(n,o["基期营业总收入"]); ppe=np.full(n,o["固定资产净额"])
    nwc_prev=np.full(n, o["应收票据及应收账款"]+o["存货"]+o["其他流动资产"]
                        -o["应付票据及应付账款"]-o["其他流动负债"])
    MY=D["期中折现开关"]; pv=np.zeros(n); fcff1=None
    for t in range(ny):
        if anchored and t==0:
            rev=np.full(n, o["基期营业总收入"]*(1+g_live[0]))
            cogs=rev*(1-m_live[0])
        else:
            rev=rev*(1+g_live[t]+dg+eg[t])
            cogs=rev*(1-(m_live[t]+dm+em[t]))
        ebit=rev-cogs-rev*opex+rev*oth
        dep=ppe*G["折旧率（占期初固定资产净额）"]
        amt=rev*G["无形摊销 / 营收"]      # 折旧/摊销拆分后 MC 一直漏了摊销，与确定性模型不一致
        fcx=(F["ucx"]/F["by"]) if (F["by"]>0 and F["yr"]<=t+1<=F["yr"]+F["by"]-1) else 0.0
        capex=rev*np.clip(cx_live[t]+dcx,0.0,None)+fcx
        ncx=-rev*G["其他经营损益 / 营收"]      # 非现金加回（与确定性模型一致）
        ppe=ppe+capex*(1-CXI)-dep-ncx*IPS  # 无形资本开支不进固定资产；减值按比例减记
        nwc=(rev*G["应收周转天数 DSO"]/365 + cogs*G["存货周转天数 DIO"]/365
             + rev*G["其他流动资产 / 营收"] - cogs*G["应付周转天数 DPO"]/365
             - rev*G["其他流动负债 / 营收"])
        fcff=ebit*(1-tax)+dep+amt-capex-(nwc-nwc_prev)+ncx; nwc_prev=nwc
        if t==0: fcff1=fcff.copy()
        pv+=fcff/(1+W)**(t+1-0.5*MY)
    tt=D["终值所得税率 t_term"]
    nopat_T=(ebit+ncx)*(1-tt)          # 正常化终值 EBIT
    roic=nopat_T/(nwc+ppe)
    rr=np.where(roic>0, gt/np.where(roic>0,roic,1.0), np.inf)
    # ROIC ≤ WACC 时增长不创造价值，Gordon 式 (1−g/ROIC) 会失真甚至变负。
    # 经济上正确的处理是降级为【无增长永续】TV = NOPAT / WACC。
    degrade = (roic<=W) | (rr>=1) | (rr<0) | ~np.isfinite(rr)
    tf_gordon=nopat_T*(1+gt)*(1-np.where(np.isfinite(rr),rr,0.0))
    tv=np.where(degrade, nopat_T/W, tf_gordon/(W-gt))
    neg_ebit=float((ebit<=0).mean()); degraded=float(degrade.mean())
    pvtv=tv/(1+W)**(ny-0.5*MY)
    ev0=pv+pvtv
    rl=D["估值基准日滚动年数"]
    ev=ev0*(1+W)**rl - fcff1*rl*(1+W)**(rl/2)
    nd=(o["短期借款"]+o["一年内到期非流动负债"]+o["长期借款"]-o["货币资金"]-o["交易性金融资产"])
    gx=ev-nd+D["非经营资产（加回）"]+F["net"]/(1+W)**F["disc"]
    mi=np.where(D["少数股东扣除方式"]==1, o["少数股东权益"], gx*D["少数股东利润占比"])
    ps=(gx-mi)/(o["总股本（亿股）"]+F["ns"])
    return dict(ps=ps, ev=ev, roic=roic, wacc=W, gt=gt, dg=dg, dm=dm, beta=beta,
                tax=tax, dcx=dcx, clipped=clipped, rej_detail=rej_detail,
                degraded=degraded, neg_ebit=neg_ebit,
                se=float(ps.std()/math.sqrt(n)),
                terminal_share=float(np.median(pvtv/ev0)))


def tornado(inp, P):
    """龙卷风：逐变量取其分布的 P10/P90（其余不随机），量测每股摆幅。"""
    from scipy.stats import norm
    zero={k:(0.0 if str(k).startswith("sd_") else v) for k,v in P.items()}
    zero["n_paths"]=1
    base=float(simulate(inp, zero)["ps"][0])
    labels={"sd_g":"营收增速（路径平移）","sd_m":"毛利率（路径平移）","sd_beta":"Beta → WACC",
            "sd_gt":"永续增长率 g","sd_cx":"资本开支率","sd_tax":"实际税率"}
    rows=[]
    for key,lab in labels.items():
        sd=float(P.get(key,0.0) or 0.0)
        if sd<=0: continue
        lo=_shifted(inp, zero, key, norm.ppf(0.10)*sd)
        hi=_shifted(inp, zero, key, norm.ppf(0.90)*sd)
        rows.append(dict(name=lab, low=min(lo,hi), high=max(lo,hi), swing=abs(hi-lo)))
    rows.sort(key=lambda x:-x["swing"])
    return base, rows


def _shifted(inp, P, key, shift):
    """把某个变量的均值平移 shift，其余不随机，返回单路径每股价值。"""
    import copy
    inp2=copy.deepcopy(inp); P2=dict(P); P2["n_paths"]=1
    if key=="sd_g":     inp2["growth"]=[x+shift for x in inp["growth"]]

    elif key=="sd_m":   inp2["gm"]=[x+shift for x in inp["gm"]]

    elif key=="sd_beta":inp2["dcf"]=dict(inp["dcf"]); inp2["dcf"]["Beta（5年"]+=shift
    elif key=="sd_gt":  inp2["dcf"]=dict(inp["dcf"]); inp2["dcf"]["永续增长率 g"]+=shift
    elif key=="sd_cx":  inp2["capex"]=[x+shift for x in inp["capex"]]
    elif key=="sd_tax": inp2["gen"]=dict(inp["gen"]); inp2["gen"]["实际所得税率"]+=shift
    return float(simulate(inp2, P2)["ps"][0])


# ---------- Excel 写回 ----------
def _sec(ws,r,t,span=8):
    ws.cell(r,1,t).font=F_HDR
    for c in range(1,span+1): ws.cell(r,c).fill=P_N
    return r+1

def build_sheet(wb, inp, params, hist_note, res, base_ps, torn):
    if SHEET in wb.sheetnames: del wb[SHEET]
    ws=wb.create_sheet(SHEET); ws.sheet_properties.tabColor="7030A0"
    ws.column_dimensions["A"].width=34
    for c in "BCDEFGH": ws.column_dimensions[c].width=14
    ws.column_dimensions["I"].width=54
    ws["A1"]="蒙特卡洛模拟"; ws["A1"].font=Font(bold=True,size=15,color=NAVY)
    ws["A2"]=("计算在 Python（向量化 numpy），参数在本页可编辑。改完参数后重跑："
              "python3 ashare_mc.py <本文件路径>")
    ws["A2"].font=Font(size=9,italic=True)
    ws["A3"]="固定随机种子 → 结果完全可复现。所有标准差置 0 时必须复现确定性基准（见下方自检行）。"
    ws["A3"].font=Font(size=9,color="C00000",bold=True)
    r=5; r=_sec(ws,r,"一、模拟参数（蓝色可编辑）")
    prow={}
    for k,lab,_,fmt,note in PARAM_SPEC:
        ws.cell(r,1,lab).font=Font(size=10); ws.cell(r,1).border=BX
        cc=ws.cell(r,2,params[k]); cc.font=F_IN; cc.number_format=fmt; cc.fill=P_Y
        nt=note
        if k in hist_note: nt=hist_note[k]+("｜"+note if note else "")
        ws.cell(r,9,nt).font=Font(size=9,color="808080")
        prow[k]=r; r+=1
    r+=1
    r=_sec(ws,r,"二、模拟结果")
    ps=res["ps"]; price=inp["o"]["现价（元）"]
    stats=[("路径数",int(params["n_paths"]),INT),("随机种子",int(params["seed"]),"0"),
           ("确定性基准（所有σ=0）",base_ps,CNY),
           ("★ 自检：基准 − 5_DCF 每股",None,"0.00000000"),
           ("均值",float(ps.mean()),CNY),("中位数 P50",float(np.median(ps)),CNY),
           ("标准差",float(ps.std()),CNY),
           ("P5",float(np.percentile(ps,5)),CNY),("P10",float(np.percentile(ps,10)),CNY),
           ("P25",float(np.percentile(ps,25)),CNY),("P75",float(np.percentile(ps,75)),CNY),
           ("P90",float(np.percentile(ps,90)),CNY),("P95",float(np.percentile(ps,95)),CNY),
           ("最小值",float(ps.min()),CNY),("最大值",float(ps.max()),CNY),
           ("现价",price,CNY),
           ("P（内在价值 > 现价）",float((ps>price).mean()),PCT),
           ("P（内在价值 < 现价一半）",float((ps<price*0.5).mean()),PCT),
           ("均值的标准误 SE",res["se"],CNY),
           ("均值 95% 置信区间 ±",1.96*res["se"],CNY),
           ("最大单变量拒绝率",res["clipped"],PCT),
           ("　· 增速与毛利率",res["rej_detail"]["增速与毛利率"],PCT),
           ("　· Beta",res["rej_detail"]["Beta"],PCT),
           ("　· 实际税率",res["rej_detail"]["实际税率"],PCT),
           ("　· 资本开支",res["rej_detail"]["资本开支"],PCT),
           ("　· 永续 g 被 WACC 间距压低",res["rej_detail"]["永续g"],PCT),
           ("⚠ 终值降级为无增长永续的路径",res["degraded"],PCT),
           ("⚠ 终值年 EBIT ≤ 0 的路径",res["neg_ebit"],PCT),
           ("终值占 EV 中位数",res["terminal_share"],PCT)]
    srow={}
    for lab,v,fmt in stats:
        ws.cell(r,1,lab).font=Font(bold=("★" in lab or "中位" in lab),size=10); ws.cell(r,1).border=BX
        cc=ws.cell(r,2,v); cc.number_format=fmt; cc.font=F_FX
        if "★" in lab: cc.fill=P_O
        srow[lab]=r; r+=1
    ws.cell(srow["★ 自检：基准 − 5_DCF 每股"],2).value=f"=B{srow['确定性基准（所有σ=0）']}-'5_DCF'!B{_ps_row(wb)}"
    ws.cell(srow["★ 自检：基准 − 5_DCF 每股"],9,"必须为 0：关掉全部随机后蒙特卡洛引擎须与主模型完全一致").font=Font(size=9,color="006400")
    r+=1
    r=_sec(ws,r,"三、龙卷风：各变量取 P10 / P90 时的每股价值（其余不随机）")
    ws.cell(r,1,"变量").font=F_COL; ws.cell(r,1).fill=P_B
    for i,h in enumerate(["低情形 P10","高情形 P90","摆幅"]):
        cc=ws.cell(r,2+i,h); cc.font=F_COL; cc.fill=P_B
    r+=1; t0=r
    for t in torn:
        ws.cell(r,1,t["name"]).font=Font(size=10)
        for i,v in enumerate([t["low"],t["high"],t["swing"]]):
            cc=ws.cell(r,2+i,round(v,4)); cc.number_format=CNY; cc.font=F_FX
        r+=1
    t1=r-1; r+=1
    r=_sec(ws,r,"四、分布数据（供图表引用）")
    hist,edges=np.histogram(ps, bins=40)
    ws.cell(r,1,"区间中点").font=F_COL; ws.cell(r,2,"频数").font=F_COL
    ws.cell(r,4,"分位数").font=F_COL; ws.cell(r,5,"每股价值").font=F_COL
    for c in (1,2,4,5): ws.cell(r,c).fill=P_B
    r+=1; h0=r
    qs=np.arange(1,100,2.5)
    for i in range(40):
        ws.cell(r+i,1,round(float((edges[i]+edges[i+1])/2),3)).number_format="0.0"
        ws.cell(r+i,2,int(hist[i]))
        if i < len(qs):
            ws.cell(r+i,4,round(float(qs[i]),1)).number_format="0.0"
            ws.cell(r+i,5,round(float(np.percentile(ps,qs[i])),3)).number_format=CNY
    h1=r+39
    _mc_charts(ws, h0, h1, t0, t1, h1+3)
    return ws

def _ps_row(wb):
    d=wb["5_DCF"]
    for r in range(1,d.max_row+1):
        v=d.cell(r,1).value
        if isinstance(v,str) and v.strip().startswith("每股内在价值"): return r
    return 1


def _mc_charts(ws, h0, h1, t0, t1, anchor_row):
    from ashare_build import _skin, _bar_series, _line_series, C_RED, C_NAVY, C_SLATE
    from openpyxl.chart import BarChart, LineChart, Reference
    # ① 分布直方图
    ch=BarChart(); ch.type="col"; ch.gapWidth=8
    ch.series.append(_bar_series(Reference(ws,min_col=2,min_row=h0,max_row=h1),"频数",C_RED))
    ch.set_categories(Reference(ws,min_col=1,min_row=h0,max_row=h1))
    _skin(ch,"每股内在价值的模拟分布",h=8.4,w_=17.5)
    ws.add_chart(ch,f"A{anchor_row}")
    # ② 分位数曲线（CDF）
    ch=LineChart()
    ch.series.append(_line_series(Reference(ws,min_col=5,min_row=h0,max_row=h0+39),"每股内在价值",C_NAVY,marker=False))
    ch.set_categories(Reference(ws,min_col=4,min_row=h0,max_row=h0+39))
    _skin(ch,"累积分布：分位数 → 每股内在价值",h=8.4,w_=17.5)
    ws.add_chart(ch,f"K{anchor_row}")
    # ③ 龙卷风
    ch=BarChart(); ch.type="bar"; ch.gapWidth=45
    ch.series.append(_bar_series(Reference(ws,min_col=4,min_row=t0,max_row=t1),"每股价值摆幅",C_SLATE))
    ch.set_categories(Reference(ws,min_col=1,min_row=t0,max_row=t1))
    _skin(ch,"龙卷风：各变量对估值的影响幅度（P10 → P90）",labels=True,label_fmt="0.0",h=8.4,w_=17.5)
    ws.add_chart(ch,f"A{anchor_row+18}")


def read_params(ws) -> dict | None:
    if ws is None: return None
    rm=_rowmap(ws); out={}
    for k,lab,dflt,_,_ in PARAM_SPEC:
        r=rm.get(lab)
        if r is None: return None
        v=ws.cell(r,2).value
        try: out[k]=float(v)
        except (TypeError,ValueError): out[k]=dflt
    return out


def run(path: str, hist: dict | None = None, n: int | None = None,
        seed: int | None = None, quiet: bool = False) -> dict:
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import ashare_verify as V
    inp=V.read_inputs(path)
    wb=load_workbook(path)
    P=read_params(wb[SHEET]) if SHEET in wb.sheetnames else None
    src="沿用文件内已有参数" if P else "首次生成，使用默认参数"
    if P is None:
        P=defaults(hist or {})
    if n:    P["n_paths"]=n
    if seed: P["seed"]=seed
    hn={}
    if hist:
        hn={"sd_g":f"历史增速标准差 {hist.get('sd_growth',0)*100:.1f}pp",
            "sd_m":f"历史毛利率标准差 {hist.get('sd_margin',0)*100:.1f}pp",
            "rho":f"历史相关系数 {hist.get('rho',0):+.3f}（{hist.get('rho_n',0)} 个年度观测）"}
    res=simulate(inp, P)
    zero={k:(0.0 if str(k).startswith("sd_") else v) for k,v in P.items()}; zero["n_paths"]=1
    base_ps=float(simulate(inp, zero)["ps"][0])
    base, torn = tornado(inp, P)
    # 主动断言：关掉全部随机后，MC 引擎必须与确定性模型完全一致
    _rows,_dv=V.recompute(inp)
    _gap=abs(base_ps-_dv["ps"])
    if _gap>1e-6:
        raise AssertionError(f"蒙特卡洛引擎与确定性模型不一致：MC基准 {base_ps:.8f} vs 模型 {_dv['ps']:.8f}（差 {_gap:.2e}）")
    build_sheet(wb, inp, P, hn, res, base_ps, torn)
    wb.save(path)
    ps=res["ps"]; price=inp["o"]["现价（元）"]
    out=dict(n=int(P["n_paths"]), seed=int(P["seed"]), base=base_ps,
             mean=float(ps.mean()), p50=float(np.median(ps)), sd=float(ps.std()),
             p5=float(np.percentile(ps,5)), p95=float(np.percentile(ps,95)),
             p_above=float((ps>price).mean()), clipped=res["clipped"],
             se=res["se"], degraded=res["degraded"], neg_ebit=res["neg_ebit"],
             price=price, tornado=torn, src=src)
    if not quiet:
        print(f"   蒙特卡洛（{src}）：{out['n']:,} 路径，种子 {out['seed']}")
        print(f"     确定性基准 ¥{out['base']:.2f}  |  中位 ¥{out['p50']:.2f}  均值 ¥{out['mean']:.2f}  标准差 ¥{out['sd']:.2f}")
        print(f"     P5 ¥{out['p5']:.2f}  P95 ¥{out['p95']:.2f}  |  P(内在价值>现价 ¥{price}) = {out['p_above']*100:.2f}%")
        print(f"     均值标准误 ±¥{out['se']:.3f}（95%CI ±¥{1.96*out['se']:.3f}）")
        if out["clipped"]>0.02:
            d=res["rej_detail"]; top=max(d,key=d.get)
            print(f"     ⚠ 拒绝重抽率最高的变量：{top} {d[top]*100:.1f}%（边界偏窄）")
        if out["degraded"]>0.001:
            print(f"     ⚠ {out['degraded']*100:.2f}% 路径终值降级为无增长永续（ROIC ≤ WACC 或 g/ROIC ≥ 1）")
        if out["neg_ebit"]>0.0005:
            print(f"     ⚠ {out['neg_ebit']*100:.2f}% 路径终值年 EBIT ≤ 0")
        print(f"     ✓ 引擎自检：关闭随机后与确定性模型一致（差 {_gap:.1e}）")
        print(f"     龙卷风前三：" + "、".join(f"{t['name']} ±{t['swing']/2:.1f}元" for t in torn[:3]))
    return out


# ---------- 截尾抽样：拒绝重抽（避免 np.clip 造成边界点质量）----------
def _rejection(rng, mu, sd, lo, hi, n, max_iter=60):
    if not sd: return np.full(n, mu), 0.0
    out=np.empty(n); filled=0; tried=0; rejected=0
    while filled<n and tried<max_iter:
        m=int((n-filled)*1.6)+64
        x=rng.normal(mu, sd, m)
        ok=(x>=lo)&(x<=hi); rejected+=int((~ok).sum())
        take=x[ok][:n-filled]
        out[filled:filled+len(take)]=take; filled+=len(take); tried+=1
    if filled<n: out[filled:]=np.clip(rng.normal(mu,sd,n-filled), lo, hi)
    total=n+rejected
    return out, float(rejected/total) if total else 0.0

def _rejection_pair(rng, L, sd1, sd2, lo1, hi1, lo2, hi2, n, max_iter=60):
    """成对拒绝重抽：保持 Cholesky 注入的相关性，同时满足各自的截尾区间。"""
    if not sd1 and not sd2:
        return np.zeros(n), np.zeros(n), 0.0
    a=np.empty(n); b=np.empty(n); filled=0; tried=0; rejected=0
    while filled<n and tried<max_iter:
        m=int((n-filled)*1.8)+64
        zc=L@rng.standard_normal((2,m))
        x=zc[0]*sd1; y=zc[1]*sd2
        ok=(x>=lo1)&(x<=hi1)&(y>=lo2)&(y<=hi2)
        rejected+=int((~ok).sum())
        k=min(int(ok.sum()), n-filled)
        a[filled:filled+k]=x[ok][:k]; b[filled:filled+k]=y[ok][:k]
        filled+=k; tried+=1
    if filled<n:
        zc=L@rng.standard_normal((2,n-filled))
        a[filled:]=np.clip(zc[0]*sd1,lo1,hi1); b[filled:]=np.clip(zc[1]*sd2,lo2,hi2)
    total=n+rejected
    return a, b, float(rejected/total) if total else 0.0

def _corr_noise(rng, L, sd1, sd2, ny, n):
    """逐年扰动也注入同一相关结构——历史 ρ 度量的就是年度间的同向性。"""
    if not sd1 and not sd2:
        return np.zeros((ny,n)), np.zeros((ny,n))
    zc=L@rng.standard_normal((2, ny*n))
    return (zc[0]*sd1).reshape(ny,n), (zc[1]*sd2).reshape(ny,n)


if __name__ == "__main__":
    import argparse
    ap=argparse.ArgumentParser(description="对已生成的模型重跑蒙特卡洛（读取文件内参数区）")
    ap.add_argument("xlsx"); ap.add_argument("--n",type=int); ap.add_argument("--seed",type=int)
    a=ap.parse_args()
    if not os.path.exists(a.xlsx): print(f"❌ 文件不存在：{a.xlsx}",file=sys.stderr); sys.exit(2)
    run(a.xlsx, n=a.n, seed=a.seed)
    print(f"\n✅ 已写回 {a.xlsx}")

